import base64
<<<<<<< HEAD
import argparse
import colorsys
import json
import mimetypes
=======
import json
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
import os
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
import traceback
import uuid
import zipfile
<<<<<<< HEAD
from copy import copy
from io import BytesIO
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse
from urllib.error import HTTPError, URLError
from urllib.request import Request as UrlRequest, urlopen

from openpyxl import Workbook, load_workbook
<<<<<<< HEAD
from openpyxl.utils import get_column_letter
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
from PIL import Image

try:
    import requests
except Exception:
    requests = None

try:
    from fastapi import FastAPI, Request as FastAPIRequest
    from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
    import uvicorn

    FASTAPI_AVAILABLE = True
except Exception:
    FastAPI = None
    FastAPIRequest = None
    FileResponse = None
    HTMLResponse = None
    JSONResponse = None
    uvicorn = None
    FASTAPI_AVAILABLE = False


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "tasks.db"
<<<<<<< HEAD
CONFIG_PATH = BASE_DIR / "app_config.json"
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
OUTPUT_DIR = BASE_DIR / "output"
IMPORT_DIR = BASE_DIR / "imports"
EXPORT_DIR = BASE_DIR / "exports"
LOG_DIR = BASE_DIR / "logs"
STATIC_DIR = BASE_DIR / "static"
<<<<<<< HEAD
EXCEL_ONLY_MODE = True
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e

STATUSES = {
    "pending",
    "generating",
    "collaging",
    "collaging_failed",
    "quality_failed",
    "waiting_confirm",
    "awaiting_confirm",
<<<<<<< HEAD
    "detail_generating",
    "detail_done",
    "detail_failed",
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    "approved",
    "rejected",
    "failed",
    "quota_exhausted",
}

STATUS_LABELS = {
    "pending": "待处理",
    "generating": "生成中",
    "collaging": "拼图中",
    "collaging_failed": "拼图失败",
    "quality_failed": "质检失败",
    "waiting_confirm": "待确认",
    "awaiting_confirm": "待确认",
<<<<<<< HEAD
    "detail_generating": "详情页生成中",
    "detail_done": "详情页完成",
    "detail_failed": "详情页失败",
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    "approved": "已通过",
    "rejected": "已驳回",
    "failed": "真失败",
    "quota_exhausted": "额度耗尽",
}

IMAGE_DELAY_SECONDS = 3
IMAGE_TIMEOUT_SECONDS = 900
<<<<<<< HEAD
HTTP_IMAGE_TIMEOUT_SECONDS = 300
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
LOCAL_IMAGE_PATHS = (
    "/v1/images/generations",
    "/api/images/generations",
    "/images/generations",
)
queue_lock = threading.Lock()
worker_lock = threading.Lock()
worker_running = False
queue_control = {"state": "idle", "pause_requested": False, "stop_requested": False}
<<<<<<< HEAD
detail_lock = threading.Lock()
detail_worker_running = False
detail_queue = []
detail_progress_state = {"state": "idle", "completed": 0, "total": 0, "message": ""}
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
quota_state = {"status": "unknown", "message": "尚未检测", "updated_at": None}
api_config_state = {"authenticated": False, "message": "尚未配置本地服务地址", "updated_at": None}


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def ensure_admin():
    if os.name != "nt":
        return
    try:
        import ctypes

        if ctypes.windll.shell32.IsUserAnAdmin():
            return
        params = subprocess.list2cmdline([str(Path(__file__).resolve())])
        result = ctypes.windll.shell32.ShellExecuteW(
            None,
            "runas",
            sys.executable,
            params,
            str(BASE_DIR),
            1,
        )
        if result <= 32:
            print(f"WARNING: 管理员权限申请失败，ShellExecuteW 返回：{result}；继续以当前权限运行", flush=True)
            return
        sys.exit(0)
    except Exception as exc:
        print(f"WARNING: 管理员权限自动申请失败：{exc}；继续以当前权限运行", flush=True)


def ensure_dirs():
    for path in (OUTPUT_DIR, IMPORT_DIR, EXPORT_DIR, LOG_DIR, STATIC_DIR):
        path.mkdir(parents=True, exist_ok=True)


def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    ensure_dirs()
<<<<<<< HEAD
    if EXCEL_ONLY_MODE:
        print("Excel-only 模式：任务状态直接读写 Excel，不再初始化 tasks.db", flush=True)
        return
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    with connect() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sku TEXT NOT NULL UNIQUE,
                product_name TEXT,
                prompt TEXT,
                raw_json TEXT NOT NULL DEFAULT '{}',
                status TEXT NOT NULL DEFAULT 'pending',
                image_path TEXT,
                spu TEXT,
                color_code TEXT,
                color_name TEXT,
                image_folder TEXT,
                image_name TEXT,
                add_logo TEXT,
                logo_path TEXT,
                brand_check TEXT,
<<<<<<< HEAD
                reference_image_path TEXT,
                scene_description TEXT,
                detail_style TEXT,
                detail_prompt TEXT,
                detail_count INTEGER,
                detail_status TEXT,
                detail_error TEXT,
                detail_generated_at TEXT,
                detail_progress TEXT,
                ai_title TEXT,
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                optimized_prompt TEXT,
                prompt_snapshot TEXT,
                generated_at TEXT,
                generation_progress TEXT,
                error_message TEXT,
                retry_count INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        existing_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(tasks)").fetchall()
        }
        for column, ddl in {
            "spu": "ALTER TABLE tasks ADD COLUMN spu TEXT",
            "color_code": "ALTER TABLE tasks ADD COLUMN color_code TEXT",
            "color_name": "ALTER TABLE tasks ADD COLUMN color_name TEXT",
            "image_folder": "ALTER TABLE tasks ADD COLUMN image_folder TEXT",
            "image_name": "ALTER TABLE tasks ADD COLUMN image_name TEXT",
            "add_logo": "ALTER TABLE tasks ADD COLUMN add_logo TEXT",
            "logo_path": "ALTER TABLE tasks ADD COLUMN logo_path TEXT",
            "brand_check": "ALTER TABLE tasks ADD COLUMN brand_check TEXT",
<<<<<<< HEAD
            "reference_image_path": "ALTER TABLE tasks ADD COLUMN reference_image_path TEXT",
            "scene_description": "ALTER TABLE tasks ADD COLUMN scene_description TEXT",
            "detail_style": "ALTER TABLE tasks ADD COLUMN detail_style TEXT",
            "detail_prompt": "ALTER TABLE tasks ADD COLUMN detail_prompt TEXT",
            "detail_count": "ALTER TABLE tasks ADD COLUMN detail_count INTEGER",
            "detail_status": "ALTER TABLE tasks ADD COLUMN detail_status TEXT",
            "detail_error": "ALTER TABLE tasks ADD COLUMN detail_error TEXT",
            "detail_generated_at": "ALTER TABLE tasks ADD COLUMN detail_generated_at TEXT",
            "detail_progress": "ALTER TABLE tasks ADD COLUMN detail_progress TEXT",
            "ai_title": "ALTER TABLE tasks ADD COLUMN ai_title TEXT",
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            "optimized_prompt": "ALTER TABLE tasks ADD COLUMN optimized_prompt TEXT",
            "prompt_snapshot": "ALTER TABLE tasks ADD COLUMN prompt_snapshot TEXT",
            "generated_at": "ALTER TABLE tasks ADD COLUMN generated_at TEXT",
            "generation_progress": "ALTER TABLE tasks ADD COLUMN generation_progress TEXT",
        }.items():
            if column not in existing_columns:
                conn.execute(ddl)
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER,
                sku TEXT,
                event_type TEXT NOT NULL,
                message TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS config (
                key TEXT PRIMARY KEY,
                value TEXT,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
            """
        )


def log_event(conn, task_id, sku, event_type, message=""):
    conn.execute(
        "INSERT INTO events(task_id, sku, event_type, message, created_at) VALUES (?, ?, ?, ?, ?)",
        (task_id, sku, event_type, message, now_iso()),
    )


def normalize_row(headers, values):
    data = {}
    for idx, header in enumerate(headers):
        if header:
            data[str(header).strip()] = "" if values[idx] is None else str(values[idx]).strip()
    return data


def normalize_header_name(value):
<<<<<<< HEAD
    return str(value or "").replace("\u00a0", " ").strip().replace("：", ":").rstrip(":").lower()


def normalize_excel_column_name(value):
    return str(value or "").replace("\u00a0", " ").strip()
=======
    return str(value or "").strip().replace("：", ":").rstrip(":").lower()
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e


def pick_field(data, candidates):
    lower_map = {normalize_header_name(k): v for k, v in data.items()}
    for key in candidates:
        if key in data and data[key]:
            return data[key]
        value = lower_map.get(normalize_header_name(key))
        if value:
            return value
    return ""


<<<<<<< HEAD
def parse_int_field(value, default=0, minimum=None, maximum=None):
    text = str(value or "").strip()
    if not text:
        return default
    try:
        number = int(float(text))
    except Exception:
        return default
    if minimum is not None:
        number = max(minimum, number)
    if maximum is not None:
        number = min(maximum, number)
    return number


=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
def strip_image_name_suffix(value, image_name):
    text = (value or "").strip()
    suffix = (image_name or "").strip()
    if not text or not suffix:
        return text
    safe_text = safe_path_part(text, text)
    safe_suffix = safe_path_part(suffix, suffix)
    candidates = [
        f"_{suffix}",
        f"/{suffix}",
        f"_{safe_suffix}",
        f"/{safe_suffix}",
    ]
    for candidate in candidates:
        if text.endswith(candidate):
            return text[: -len(candidate)].rstrip("_/ ")
        if safe_text.endswith(candidate):
            return safe_text[: -len(candidate)].rstrip("_/ ")
    return text


<<<<<<< HEAD
def split_list_text(value):
    text = (value or "").strip()
    if not text:
        return []
    for sep in ("，", "、", "；", ";", "|", "\n", "\r"):
        text = text.replace(sep, "/")
    return [part.strip() for part in text.split("/") if part.strip()]


def find_header_row(rows, required_candidates):
    required = {normalize_header_name(name) for name in required_candidates}
    for index, row in enumerate(rows[:20]):
        normalized = {normalize_header_name(cell) for cell in row if cell is not None}
        if normalized & required:
            return index
    return 0


def select_import_sheet(wb, table_mode="auto"):
    if table_mode == "old":
        for name in ("任务表", "任务", wb.sheetnames[0]):
            if name in wb.sheetnames:
                return wb[name], "old"
    if table_mode == "new":
        for name in ("商品总表", "Sheet1", wb.sheetnames[0]):
            if name in wb.sheetnames:
                return wb[name], "new"
    if "商品总表" in wb.sheetnames:
        return wb["商品总表"], "new"
    if "任务表" in wb.sheetnames:
        return wb["任务表"], "old"
    return wb[wb.sheetnames[0]], "old"


def import_excel(path, table_mode="auto"):
    wb = load_workbook(path)
    ws, resolved_mode = select_import_sheet(wb, table_mode)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"imported": 0, "updated": 0, "skipped": 0, "mode": resolved_mode}

    header_index = find_header_row(rows, ["SKU", "SPU", "图片生成提示词"])
    headers = ["" if cell is None else str(cell).strip() for cell in rows[header_index]]
    data_rows = rows[header_index + 1 :]
    stats = {"imported": 0, "updated": 0, "skipped": 0, "mode": resolved_mode, "sheet": ws.title}
    set_config("last_import_path", str(path.resolve()))
    set_config("last_import_table_mode", resolved_mode)
    set_config("last_import_sheet", ws.title)
=======
def import_excel(path):
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"imported": 0, "updated": 0, "skipped": 0}

    headers = ["" if cell is None else str(cell).strip() for cell in rows[0]]
    stats = {"imported": 0, "updated": 0, "skipped": 0}
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    spu_color_counters = {}
    with connect() as conn:
        for row in conn.execute("SELECT spu, image_name FROM tasks WHERE spu IS NOT NULL AND image_name IS NOT NULL").fetchall():
            image_name = row["image_name"] or ""
            prefix = image_name.split("_", 1)[0]
            if prefix.isdigit():
                spu_color_counters[row["spu"]] = max(spu_color_counters.get(row["spu"], 0), int(prefix))
<<<<<<< HEAD
        for values in data_rows:
            data = normalize_row(headers, list(values) + [None] * (len(headers) - len(values)))
            if resolved_mode == "new":
                excel_sku = pick_field(data, ["SKU", "sku", "货号", "商品编码"])
                raw_spu = excel_sku
            else:
                raw_spu = pick_field(data, ["SPU", "spu", "产品系列编号", "系列编号", "产品系列", "款号"])
                excel_sku = pick_field(data, ["SKU", "sku", "货号", "商品编码"])
=======
        for values in rows[1:]:
            data = normalize_row(headers, list(values) + [None] * (len(headers) - len(values)))
            raw_spu = pick_field(data, ["SPU", "spu", "产品系列编号", "系列编号", "产品系列", "款号"])
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            color_name = pick_field(data, ["图片颜色", "颜色名称", "颜色", "Color", "color_name"])
            image_name = pick_field(data, ["图片命名", "image_name"])
            spu = strip_image_name_suffix(raw_spu, image_name)
            image_name_source = "excel" if image_name else "empty"
            if not image_name and color_name:
                next_index = spu_color_counters.get(spu, 0) + 1
                spu_color_counters[spu] = next_index
                image_name = f"{next_index:02d}_{color_name}"
                image_name_source = "auto_color"
                spu = strip_image_name_suffix(raw_spu, image_name)
<<<<<<< HEAD
            sku = spu or image_name or excel_sku
            sku_source = "new_sku" if resolved_mode == "new" and sku else ("spu" if spu else ("image_name" if image_name else "excel_sku"))
=======
            excel_sku = pick_field(data, ["SKU", "sku", "货号", "商品编码"])
            sku = spu or image_name or excel_sku
            sku_source = "spu" if spu else ("image_name" if image_name else "excel_sku")
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            if not sku:
                stats["skipped"] += 1
                continue
            print(
                f"DEBUG import row: sku={sku}, spu={spu}, image_name={image_name}, excel_sku={excel_sku}, sku_source={sku_source}",
                flush=True,
            )
            print(
                "DEBUG import row detail:",
                f"sku_repr={sku!r}",
                f"sku_source={sku_source}",
                f"raw_spu_repr={raw_spu!r}",
                f"spu_repr={spu!r}",
                f"excel_sku_repr={excel_sku!r}",
                f"color_name_repr={color_name!r}",
                f"image_name_repr={image_name!r}",
                f"image_name_source={image_name_source}",
                flush=True,
            )
            product_name = pick_field(data, ["商品名称", "产品名称", "标题", "品名", "Name"])
            prompt = pick_field(data, ["图片生成提示词", "生图提示词", "Prompt", "prompt", "提示词"])
            color_code = pick_field(data, ["颜色编号", "色号", "Color Code", "color_code"])
            if not color_code and image_name and "_" in image_name:
                color_code = image_name.split("_", 1)[0]
            add_logo = pick_field(data, ["是否添加品牌logo", "是否添加品牌Logo", "add_logo"])
            logo_path = pick_field(data, ["logo文件地址", "Logo文件地址", "logo_path"])
            brand_check = pick_field(data, ["需校验品牌文字/元素", "brand_check"])
<<<<<<< HEAD
            reference_image_path = pick_field(data, ["参考图地址", "reference_image_path", "参考图", "参考图片"])
            scene_description = pick_field(data, ["场景描述", "scene_description", "使用场景"])
            detail_style = pick_field(data, ["详情页风格", "detail_style", "风格"])
            detail_prompt = pick_field(data, ["详情页提示词", "detail_prompt", "详情提示词", "详情页生图提示词"])
            detail_count = parse_int_field(pick_field(data, ["详情页数量", "detail_count", "详情图数量"]), default=6, minimum=1, maximum=20)
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            raw_json = json.dumps(data, ensure_ascii=False)
            existing = conn.execute("SELECT id, sku FROM tasks WHERE sku = ?", (sku,)).fetchone()
            if not existing and spu and image_name:
                existing = conn.execute(
                    "SELECT id, sku FROM tasks WHERE spu = ? AND image_name = ?",
                    (spu, image_name),
                ).fetchone()
            if not existing and excel_sku:
                existing = conn.execute("SELECT id, sku FROM tasks WHERE sku = ?", (excel_sku,)).fetchone()
            if existing:
                conn.execute(
                    """
                    UPDATE tasks
                    SET sku = ?, product_name = ?, prompt = ?, raw_json = ?, spu = ?, color_code = ?,
                        color_name = ?, image_name = ?, add_logo = ?, logo_path = ?, brand_check = ?,
<<<<<<< HEAD
                        reference_image_path = ?, scene_description = ?, detail_style = ?,
                        detail_prompt = ?, detail_count = ?, prompt_snapshot = ?, updated_at = ?
=======
                        prompt_snapshot = ?, updated_at = ?
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                    WHERE id = ?
                    """,
                    (
                        sku,
                        product_name,
                        prompt,
                        raw_json,
                        spu,
                        color_code,
                        color_name,
                        image_name,
                        add_logo,
                        logo_path,
                        brand_check,
<<<<<<< HEAD
                        reference_image_path,
                        scene_description,
                        detail_style,
                        detail_prompt,
                        detail_count,
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                        prompt,
                        now_iso(),
                        existing["id"],
                    ),
                )
                log_event(conn, existing["id"], sku, "excel_update", f"Excel updated source fields; old_sku={existing['sku']}")
                stats["updated"] += 1
            else:
                cur = conn.execute(
                    """
                    INSERT INTO tasks(
                        sku, product_name, prompt, raw_json, spu, color_code, color_name,
<<<<<<< HEAD
                        image_name, add_logo, logo_path, brand_check, reference_image_path,
                        scene_description, detail_style, detail_prompt, detail_count, prompt_snapshot,
                        status, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
=======
                        image_name, add_logo, logo_path, brand_check, prompt_snapshot,
                        status, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                    """,
                    (
                        sku,
                        product_name,
                        prompt,
                        raw_json,
                        spu,
                        color_code,
                        color_name,
                        image_name,
                        add_logo,
                        logo_path,
                        brand_check,
<<<<<<< HEAD
                        reference_image_path,
                        scene_description,
                        detail_style,
                        detail_prompt,
                        detail_count,
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                        prompt,
                        now_iso(),
                        now_iso(),
                    ),
                )
                log_event(conn, cur.lastrowid, sku, "excel_import", "Excel imported task")
                stats["imported"] += 1
    return stats


def list_tasks(status=None, q=None):
    sql = "SELECT * FROM tasks"
    params = []
    clauses = []
    if status and status != "all":
        clauses.append("status = ?")
        params.append(status)
    if q:
        clauses.append("(sku LIKE ? OR product_name LIKE ?)")
        params.extend([f"%{q}%", f"%{q}%"])
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY updated_at DESC, id DESC"
    with connect() as conn:
        tasks = [dict(row) for row in conn.execute(sql, params).fetchall()]
    for task in tasks:
        if task["status"] in ("waiting_confirm", "awaiting_confirm") and not main_image_path_for_task(task):
            task["status"] = "collaging_failed"
            task["error_message"] = "主图文件未找到，请重新拼图"
            task["generation_progress"] = None
            with connect() as conn:
                conn.execute(
                    """
                    UPDATE tasks
                    SET status = 'collaging_failed', error_message = ?, generation_progress = NULL, updated_at = ?
                    WHERE id = ?
                    """,
                    (task["error_message"], now_iso(), task["id"]),
                )
        task["status_label"] = STATUS_LABELS.get(task["status"], task["status"])
        if task["status"] == "generating" and task.get("generation_progress"):
            task["status_label"] = task["generation_progress"]
<<<<<<< HEAD
        if task["status"] == "detail_generating" and task.get("detail_progress"):
            task["status_label"] = task["detail_progress"]
        if task["status"] == "detail_failed" and task.get("detail_error") and not task.get("error_message"):
            task["error_message"] = task["detail_error"]
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        task["raw"] = json.loads(task["raw_json"] or "{}")
    return tasks


def add_manual_task(payload):
    spu = (payload.get("spu") or payload.get("SPU") or "").strip()
    image_name = (payload.get("image_name") or payload.get("图片命名") or "").strip()
    sku_input = (payload.get("sku") or payload.get("SKU") or "").strip()
    sku = spu or sku_input or image_name
    product_name = (payload.get("product_name") or "").strip()
    prompt = (payload.get("prompt") or "").strip()
<<<<<<< HEAD
    detail_prompt = (payload.get("detail_prompt") or payload.get("详情页提示词") or "").strip()
    detail_count = parse_int_field(payload.get("detail_count") or payload.get("详情页数量"), default=6, minimum=1, maximum=20)
    if not sku or not product_name or not prompt:
        raise ValueError("SKU、商品名称、生图提示词均为必填")
    raw = {"SKU": sku, "SPU": spu, "图片命名": image_name, "商品名称": product_name, "生图提示词": prompt, "详情页提示词": detail_prompt, "详情页数量": detail_count}
    with connect() as conn:
        cur = conn.execute(
            """
            INSERT INTO tasks(sku, product_name, prompt, raw_json, spu, image_name, detail_prompt, detail_count, status, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
            """,
            (sku, product_name, prompt, json.dumps(raw, ensure_ascii=False), spu, image_name, detail_prompt, detail_count, now_iso(), now_iso()),
=======
    if not sku or not product_name or not prompt:
        raise ValueError("SKU、商品名称、生图提示词均为必填")
    raw = {"SKU": sku, "SPU": spu, "图片命名": image_name, "商品名称": product_name, "生图提示词": prompt}
    with connect() as conn:
        cur = conn.execute(
            """
            INSERT INTO tasks(sku, product_name, prompt, raw_json, spu, image_name, status, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, 'pending', ?, ?)
            """,
            (sku, product_name, prompt, json.dumps(raw, ensure_ascii=False), spu, image_name, now_iso(), now_iso()),
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        )
        log_event(conn, cur.lastrowid, sku, "manual_add", "Manual task created")


def update_status(task_id, status, message=""):
    if status not in STATUSES:
        raise ValueError(f"Unsupported status: {status}")
    with connect() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            raise ValueError("Task not found")
        conn.execute(
            "UPDATE tasks SET status = ?, error_message = ?, updated_at = ? WHERE id = ?",
            (status, message or None, now_iso(), task_id),
        )
        log_event(conn, task_id, task["sku"], "status_update", f"{status}: {message}")


def reset_for_generation(task_id):
    with connect() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            raise ValueError("Task not found")
        sku = task["sku"]
        print(f"DEBUG: generate_image called for SKU: {sku}", flush=True)
        conn.execute(
            """
            UPDATE tasks
            SET status = 'pending', error_message = NULL, updated_at = ?
            WHERE id = ?
            """,
            (now_iso(), task_id),
        )
        log_event(conn, task_id, task["sku"], "queued", "Queued for generation")


def is_quota_error(text):
    lowered = (text or "").lower()
    markers = [
        "quota",
        "rate limit",
        "rate_limit",
        "too many requests",
        "usage limit",
        "message limit",
        "try again later",
        "temporarily limited",
        "429",
        "额度",
        "频率",
        "限制",
        "稍后",
    ]
    return any(marker in lowered for marker in markers)


def safe_sku(sku):
    return "".join(ch if ch.isalnum() or ch in ("-", "_", ".") else "_" for ch in sku)[:120]


def safe_path_part(value, fallback):
    text = (value or fallback or "").strip()
    cleaned = "".join(ch if ch.isalnum() or ch in ("-", "_", ".", " ", "（", "）", "(", ")") else "_" for ch in text)
    return cleaned.strip(" .")[:120] or safe_sku(fallback or "unknown")


def folder_spu_for_task(task):
    spu = (task.get("spu") or "").strip()
    if spu:
        return spu
    sku = (task.get("sku") or "").strip()
    image_name = (task.get("image_name") or "").strip()
    if sku and image_name:
        suffixes = [
            f"_{image_name}",
            f"_{safe_path_part(image_name, image_name)}",
            f"/{image_name}",
        ]
        safe_sku_value = safe_path_part(sku, sku)
        safe_image_name = safe_path_part(image_name, image_name)
        if safe_image_name:
            suffixes.append(f"_{safe_image_name}")
        for suffix in suffixes:
            if sku.endswith(suffix):
                return sku[: -len(suffix)]
            if safe_sku_value.endswith(suffix):
                return safe_sku_value[: -len(suffix)]
    return sku or "unknown"


def folder_image_name_for_task(task):
    image_name = (task.get("image_name") or "").strip()
    if image_name:
        return image_name
    color_code = task.get("color_code") or "00"
    color_name = task.get("color_name") or "默认"
    return f"{safe_path_part(color_code, '00')}_{safe_path_part(color_name, '默认')}"


def image_folder_for_task(task):
    spu_folder = safe_path_part(folder_spu_for_task(task), task.get("sku") or "unknown")
    folder_path = OUTPUT_DIR / spu_folder
<<<<<<< HEAD
=======
    print(
        "DEBUG path:",
        f"sku={task.get('sku')}",
        f"spu={task.get('spu')}",
        f"image_name={task.get('image_name')}",
        f"folder={folder_path}",
        flush=True,
    )
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    return folder_path


def legacy_image_folder_for_task(task):
    spu_folder = safe_path_part(folder_spu_for_task(task), task.get("sku") or "unknown")
    image_folder = safe_path_part(folder_image_name_for_task(task), "00_默认")
    return OUTPUT_DIR / spu_folder / image_folder


def download_base_name_for_task(task, fallback_sku):
    if task:
        return safe_path_part(folder_spu_for_task(task), fallback_sku or "unknown")
    return safe_sku(str(fallback_sku))


def zip_filename_for_skus(skus):
    normalized_skus = [str(sku) for sku in skus if str(sku).strip()]
    if len(normalized_skus) == 1:
        task = get_task_by_sku(normalized_skus[0])
        return f"{download_base_name_for_task(task, normalized_skus[0])}.zip"
    return f"batch_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"


def build_codex_prompt(task, image_file):
    product_name = task["product_name"] or task["sku"]
    user_prompt = task["prompt"] or product_name
    return f"""
Use case: product-mockup
Asset type: ecommerce main product image
Primary request: {user_prompt}
Subject: {product_name}
Style/medium: clean commercial product photography
Composition/framing: square 1024x1024 main image, centered product, clear edges
Lighting/mood: polished studio lighting, realistic shadows
Constraints: no watermark, no extra text, no logos unless explicitly requested by the prompt
""".strip()


def encode_config_value(value):
    return base64.b64encode(value.encode("utf-8")).decode("ascii")


def decode_config_value(value):
    return base64.b64decode((value or "").encode("ascii")).decode("utf-8")


<<<<<<< HEAD
def load_config_store():
    if not CONFIG_PATH.exists():
        return {}
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config_store(data):
    ensure_dirs()
    CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def set_config(key, value):
    if EXCEL_ONLY_MODE:
        data = load_config_store()
        data[key] = value
        save_config_store(data)
        return
=======
def set_config(key, value):
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO config(key, value, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = CURRENT_TIMESTAMP
            """,
            (key, encode_config_value(value)),
        )


def get_config(key):
<<<<<<< HEAD
    if EXCEL_ONLY_MODE:
        value = load_config_store().get(key)
        if value in ("", None) and DB_PATH.exists():
            try:
                with sqlite3.connect(DB_PATH) as conn:
                    conn.row_factory = sqlite3.Row
                    row = conn.execute("SELECT value FROM config WHERE key = ?", (key,)).fetchone()
                if row and row["value"]:
                    value = decode_config_value(row["value"])
                    set_config(key, value)
            except Exception:
                value = None
        return value if value not in ("", None) else None
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    with connect() as conn:
        row = conn.execute("SELECT value FROM config WHERE key = ?", (key,)).fetchone()
    if not row or not row["value"]:
        return None
    try:
        return decode_config_value(row["value"])
    except Exception:
        return None


def normalize_api_base_url(api_base_url):
    api_base_url = (api_base_url or "").strip().rstrip("/")
    if not api_base_url:
        raise ValueError("请填写 Cockpit 本地 API 服务地址")
    return api_base_url


def display_api_base_url(api_base_url):
    if not api_base_url:
        return None
    return api_base_url.replace("http://", "").replace("https://", "")


def set_api_config_state(authenticated, message):
    api_config_state.update(
        {"authenticated": bool(authenticated), "message": message, "updated_at": now_iso()}
    )
    return dict(api_config_state)


def load_api_base_url():
    api_base_url = get_config("api_base_url")
    if api_base_url:
        set_api_config_state(True, "Codex 本地服务已就绪，可以生成图片")
    else:
        set_api_config_state(False, "请先配置 Codex 本地服务地址才能生成图片")
    return api_base_url


def api_config_user_info():
    api_base_url = get_config("api_base_url")
    access_token = get_config("access_token")
    if not api_base_url:
        return {"authenticated": False, "has_api_key": False, "has_api_base_url": False, "has_access_token": bool(access_token)}
    return {
        "authenticated": True,
        "has_api_key": False,
        "has_api_base_url": True,
        "has_access_token": bool(access_token),
        "api_base_url": api_base_url,
        "api_base_url_display": display_api_base_url(api_base_url),
        "membership_status": "订阅有效",
    }


def set_api_base_url(api_base_url, access_token=None):
    api_base_url = normalize_api_base_url(api_base_url)
    set_config("api_base_url", api_base_url)
    access_token = (access_token or "").strip()
    if access_token:
        set_config("access_token", access_token)
    set_api_config_state(True, "本地服务地址已保存")
    return {"success": True, "message": "本地服务地址已保存", "user": api_config_user_info()}


def require_api_base_url():
    api_base_url = load_api_base_url()
    if not api_base_url:
        raise ValueError("请先配置 Codex 本地服务地址")
    return api_base_url


def get_access_token():
    return get_config("access_token")


def local_api_headers():
    access_token = get_access_token()
    headers = {"Content-Type": "application/json"}
    if access_token:
        headers["Authorization"] = f"Bearer {access_token}"
    return headers


def mask_token(access_token):
    if not access_token:
        return "未配置"
    if len(access_token) <= 8:
        return access_token[:2] + "****"
    return f"{access_token[:4]}****{access_token[-4:]}"


def mask_headers(headers):
    masked = dict(headers)
    authorization = masked.get("Authorization")
    if authorization and authorization.lower().startswith("bearer "):
        masked["Authorization"] = "Bearer " + mask_token(authorization[7:])
    return masked


<<<<<<< HEAD
def mask_payload(payload):
    masked = dict(payload)
    for key in ("reference_image", "input_image", "image"):
        if isinstance(masked.get(key), dict):
            item = dict(masked[key])
            if item.get("b64_json"):
                item["b64_json"] = f"<base64 {len(item['b64_json'])} chars>"
            masked[key] = item
    if isinstance(masked.get("images"), list):
        masked_images = []
        for image in masked["images"]:
            if isinstance(image, dict):
                item = dict(image)
                if item.get("b64_json"):
                    item["b64_json"] = f"<base64 {len(item['b64_json'])} chars>"
                masked_images.append(item)
            else:
                masked_images.append(image)
        masked["images"] = masked_images
    return masked


=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
def print_local_api_config():
    api_base_url = get_config("api_base_url") or "未配置"
    access_token = get_config("access_token")
    print(f"Config api_base_url: {api_base_url}", flush=True)
    print(f"Config access_token: {mask_token(access_token)}", flush=True)


def openai_error_message(exc):
    if isinstance(exc, HTTPError):
        body = exc.read().decode("utf-8", errors="replace")
        try:
            data = json.loads(body)
            message = data.get("error", {}).get("message") or body
        except Exception:
            message = body
        return f"{exc.code} - {message}"
    if isinstance(exc, URLError):
        reason = str(exc.reason)
        if "connection refused" in reason.lower() or "actively refused" in reason.lower() or "由于目标计算机积极拒绝" in reason:
            return "无法连接到 Cockpit API 服务，请确认服务已启动且端口正确"
        return f"网络错误：{exc.reason}"
    message = str(exc)
    lowered = message.lower()
    if "connection refused" in lowered or "actively refused" in lowered or "由于目标计算机积极拒绝" in message:
        return "无法连接到 Cockpit API 服务，请确认服务已启动且端口正确"
    return message


def post_local_image_api(api_base_url, payload, headers):
    last_error = None
    for path in LOCAL_IMAGE_PATHS:
        endpoint = api_base_url.rstrip("/") + path
        print(f"DEBUG image request URL: {endpoint}", flush=True)
        print(f"DEBUG image request Headers: {json.dumps(mask_headers(headers), ensure_ascii=False)}", flush=True)
<<<<<<< HEAD
        print(f"DEBUG image request Payload: {json.dumps(mask_payload(payload), ensure_ascii=False)}", flush=True)
        try:
            if requests is not None:
                response = requests.post(endpoint, json=payload, headers=headers, timeout=HTTP_IMAGE_TIMEOUT_SECONDS)
=======
        print(f"DEBUG image request Payload: {json.dumps(payload, ensure_ascii=False)}", flush=True)
        try:
            if requests is not None:
                response = requests.post(endpoint, json=payload, headers=headers, timeout=60)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                if response.status_code == 404:
                    last_error = RuntimeError(f"404 - {response.text}")
                    print(f"DEBUG endpoint returned 404, trying next path: {path}", flush=True)
                    continue
                if response.status_code >= 400:
                    raise RuntimeError(f"{response.status_code} - {response.text}")
                return response.json()
            request = UrlRequest(
                endpoint,
                data=json.dumps(payload).encode("utf-8"),
                headers=headers,
                method="POST",
            )
            try:
<<<<<<< HEAD
                with urlopen(request, timeout=HTTP_IMAGE_TIMEOUT_SECONDS) as response:
=======
                with urlopen(request, timeout=60) as response:
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                    return json.loads(response.read().decode("utf-8"))
            except HTTPError as exc:
                if exc.code == 404:
                    last_error = exc
                    print(f"DEBUG endpoint returned 404, trying next path: {path}", flush=True)
                    continue
                raise
        except Exception as exc:
            error_text = str(exc).lower()
            if "connection refused" in error_text or "actively refused" in error_text or "由于目标计算机积极拒绝" in str(exc):
                raise RuntimeError("无法连接到 Cockpit API 服务，请确认服务已启动且端口正确") from exc
            raise
    raise RuntimeError(f"所有图片生成端点都返回 404，请确认 Cockpit API 路径是否正确：{last_error}")


<<<<<<< HEAD
def resolve_local_path(path_text):
    path = Path(path_text or "")
    if not path_text:
        return None
    if not path.is_absolute():
        path = BASE_DIR / path
    return path


def compress_reference_image(path):
    original_size = path.stat().st_size
    with Image.open(path) as image:
        image = image.convert("RGB")
        image.thumbnail((768, 768), Image.LANCZOS)

        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=75, optimize=True)
        compressed = buffer.getvalue()
        quality = 75

        if len(compressed) > 800 * 1024:
            buffer = BytesIO()
            image.save(buffer, format="JPEG", quality=60, optimize=True)
            compressed = buffer.getvalue()
            quality = 60

    print(
        "DEBUG reference image compressed:",
        f"path={path}",
        f"original={original_size} bytes",
        f"compressed={len(compressed)} bytes",
        f"quality={quality}",
        flush=True,
    )
    return compressed


def reference_payload(reference_image_path):
    path = resolve_local_path(reference_image_path)
    if not path or not path.exists():
        return None, ""
    compressed = compress_reference_image(path)
    encoded = base64.b64encode(compressed).decode("ascii")
    print(
        "DEBUG reference image payload:",
        f"filename={path.name}",
        f"base64_length={len(encoded)}",
        flush=True,
    )
    note = f"\nReference image provided locally: {path.name}. Use it only as composition, lighting, camera angle and style reference; keep the generated product faithful to the text prompt."
    return {"mime_type": "image/jpeg", "b64_json": encoded, "filename": path.name}, note


def reference_payload_variants(base_payload, reference):
    return [
        ("image", {**base_payload, "image": reference}),
        (
            "reference_image + input_image",
            {**base_payload, "reference_image": reference, "input_image": reference},
        ),
        ("images", {**base_payload, "images": [reference]}),
    ]


def call_image_api(prompt, image_file, reference_image_path=""):
    api_base_url = require_api_base_url()
    reference, reference_note = reference_payload(reference_image_path)
    base_payload = {
        "prompt": prompt + reference_note,
=======
def generate_image_with_api(prompt, image_file):
    api_base_url = require_api_base_url()
    debug_print_registered_routes()
    payload = {
        "prompt": prompt,
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        "size": "1024x1024",
        "n": 1,
    }
    headers = local_api_headers()
<<<<<<< HEAD
    if reference:
        last_reference_error = None
        for variant_name, payload in reference_payload_variants(base_payload, reference):
            print(f"DEBUG trying reference image field format: {variant_name}", flush=True)
            try:
                data = post_local_image_api(api_base_url, payload, headers)
                print(f"DEBUG reference image field format succeeded: {variant_name}", flush=True)
                break
            except Exception as exc:
                last_reference_error = exc
                print(f"DEBUG reference image field format failed: {variant_name}; error={exc}", flush=True)
        else:
            print(
                f"DEBUG all reference image payload formats failed, retrying text-only: {last_reference_error}",
                flush=True,
            )
            data = post_local_image_api(api_base_url, base_payload, headers)
    else:
        data = post_local_image_api(api_base_url, base_payload, headers)
=======
    data = post_local_image_api(api_base_url, payload, headers)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e

    item = (data.get("data") or [{}])[0]
    if item.get("b64_json"):
        image_file.write_bytes(base64.b64decode(item["b64_json"]))
        return
    if item.get("url"):
        if requests is not None:
<<<<<<< HEAD
            image_response = requests.get(item["url"], headers=headers, timeout=HTTP_IMAGE_TIMEOUT_SECONDS)
=======
            image_response = requests.get(item["url"], headers=headers, timeout=60)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            if image_response.status_code >= 400:
                raise RuntimeError(f"{image_response.status_code} - {image_response.text}")
            image_file.write_bytes(image_response.content)
        else:
            image_request = UrlRequest(item["url"], headers=headers, method="GET")
<<<<<<< HEAD
            with urlopen(image_request, timeout=HTTP_IMAGE_TIMEOUT_SECONDS) as response:
=======
            with urlopen(image_request, timeout=60) as response:
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                image_file.write_bytes(response.read())
        return
    raise RuntimeError("本地图片 API 返回为空，未找到图片数据")


<<<<<<< HEAD
def generate_image_with_api(prompt, image_file, reference_image_path=""):
    debug_print_registered_routes()
    return call_image_api(prompt, image_file, reference_image_path)


def generate_single_image(prompt, image_file, reference_image_path=""):
    return call_image_api(prompt, image_file, reference_image_path)
=======
def generate_single_image(prompt, image_file):
    api_base_url = require_api_base_url()
    payload = {
        "prompt": prompt,
        "size": "1024x1024",
        "n": 1,
    }
    headers = local_api_headers()
    data = post_local_image_api(api_base_url, payload, headers)
    item = (data.get("data") or [{}])[0]
    if item.get("b64_json"):
        image_file.write_bytes(base64.b64decode(item["b64_json"]))
        return
    if item.get("url"):
        if requests is not None:
            image_response = requests.get(item["url"], headers=headers, timeout=60)
            if image_response.status_code >= 400:
                raise RuntimeError(f"{image_response.status_code} - {image_response.text}")
            image_file.write_bytes(image_response.content)
        else:
            image_request = UrlRequest(item["url"], headers=headers, method="GET")
            with urlopen(image_request, timeout=60) as response:
                image_file.write_bytes(response.read())
        return
    raise RuntimeError("本地图片 API 返回为空，未找到图片数据")
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e


def compose_nine_grid(image_files, output_file):
    images = [Image.open(path).convert("RGB") for path in image_files]
    width, height = images[0].size
    canvas = Image.new("RGB", (width * 3, height * 3), "white")
    for index, image in enumerate(images):
        if image.size != (width, height):
            image = image.resize((width, height))
        x = (index % 3) * width
        y = (index // 3) * height
        canvas.paste(image, (x, y))
    canvas.save(output_file)
    for image in images:
        image.close()


def resolve_logo_path(logo_path):
    path = Path(logo_path)
    if not path.is_absolute():
        path = BASE_DIR / path
    return path


def add_logo_to_image(image_file, logo_path, output_file, opacity=0.82):
    resolved_logo = resolve_logo_path(logo_path)
    if not resolved_logo.exists():
        raise FileNotFoundError("Logo 文件未找到")
    base = Image.open(image_file).convert("RGBA")
    logo = Image.open(resolved_logo).convert("RGBA")
    target_width = max(base.width // 5, 1)
    ratio = target_width / logo.width
    logo = logo.resize((target_width, max(int(logo.height * ratio), 1)))
    if opacity < 1:
        alpha = logo.getchannel("A")
        alpha = alpha.point(lambda value: int(value * opacity))
        logo.putalpha(alpha)
    x = (base.width - logo.width) // 2
    y = (base.height - logo.height) // 2
    base.alpha_composite(logo, (x, y))
    base.convert("RGB").save(output_file)


def run_image_for_task(task):
    sku_dir = image_folder_for_task(task)
    sku_dir.mkdir(parents=True, exist_ok=True)
    main_file = sku_dir / "主图.png"
    logo_file = sku_dir / "主图_带logo.png"
    optimized_prompt = build_codex_prompt(task, main_file)
    prompt_snapshot = task.get("prompt") or ""
    started = now_iso()
    print(
        "DEBUG generation paths:",
        f"sku={task.get('sku')}",
        f"image_folder={sku_dir}",
        f"main_file={main_file}",
        f"logo_file={logo_file}",
        flush=True,
    )

    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = 'generating', error_message = NULL, optimized_prompt = ?,
                prompt_snapshot = ?, generation_progress = '生成中(0/9)', image_folder = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (optimized_prompt, prompt_snapshot, str(sku_dir), started, task["id"]),
        )
        log_event(conn, task["id"], task["sku"], "generation_start", "Local image API started")

    try:
        image_files = []
<<<<<<< HEAD
        colors = split_list_text(task.get("color_name"))
        for index in range(1, 10):
            image_file = sku_dir / f"{index}.png"
            color_hint = colors[index - 1] if index - 1 < len(colors) else ""
            single_prompt = (
                f"{optimized_prompt}\n"
                "White background ecommerce product image, product centered, clean cutout-style commercial lighting.\n"
                f"Image {index} of 9."
            )
            if color_hint:
                single_prompt += f"\nRequired product color/material element for this image: {color_hint}."
            else:
                single_prompt += "\nKeep the same product, vary composition subtly."
=======
        for index in range(1, 10):
            image_file = sku_dir / f"{index}.png"
            single_prompt = f"{optimized_prompt}\nVariant {index} of 9. Keep the same product and color, vary composition subtly."
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            with connect() as conn:
                conn.execute(
                    "UPDATE tasks SET generation_progress = ?, updated_at = ? WHERE id = ?",
                    (f"生成中({index}/9)", now_iso(), task["id"]),
                )
<<<<<<< HEAD
            if image_file.exists():
                print(f"DEBUG image exists, skipped: sku={task.get('sku')} file={image_file}", flush=True)
            else:
                generate_single_image(single_prompt, image_file, task.get("reference_image_path") or "")
=======
            generate_single_image(single_prompt, image_file)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            if not image_file.exists():
                raise RuntimeError(f"第 {index} 张图片未生成")
            print(f"DEBUG image saved: sku={task.get('sku')} file={image_file}", flush=True)
            image_files.append(image_file)

        with connect() as conn:
            conn.execute(
                "UPDATE tasks SET status = 'collaging', generation_progress = '拼图中', updated_at = ? WHERE id = ?",
                (now_iso(), task["id"]),
            )
        compose_nine_grid(image_files, main_file)
        print(f"DEBUG collage saved: sku={task.get('sku')} file={main_file}", flush=True)
        final_file = main_file

        add_logo = str(task.get("add_logo") or "").strip().lower()
        should_add_logo = add_logo in ("是", "yes", "y", "true", "1")
        if should_add_logo:
            logo_path = task.get("logo_path")
            if not logo_path:
                raise FileNotFoundError("Logo 文件未找到")
            add_logo_to_image(main_file, logo_path, logo_file)
            final_file = logo_file
            print(f"DEBUG logo image saved: sku={task.get('sku')} file={logo_file}", flush=True)

        if not final_file.exists():
            raise RuntimeError("九宫格主图未生成")

        with connect() as conn:
            conn.execute(
                """
                UPDATE tasks
                SET status = 'waiting_confirm', image_path = ?, image_folder = ?, error_message = NULL,
                    generation_progress = NULL, generated_at = ?, updated_at = ?,
                    retry_count = retry_count + 1
                WHERE id = ?
                """,
                (str(final_file), str(sku_dir), now_iso(), now_iso(), task["id"]),
            )
            log_event(conn, task["id"], task["sku"], "generation_success", str(final_file))
        print(
            "DEBUG db image fields updated:",
            f"sku={task.get('sku')}",
            f"image_path={final_file}",
            f"image_folder={sku_dir}",
            flush=True,
        )
        set_quota("available", "最近一次本地 API 生图成功")
        return True
    except Exception as exc:
        message = openai_error_message(exc)
        mark_generation_error(task, message)
        return False


def apply_collage_for_task(task):
    sku_dir, image_files = folder_with_complete_nine_images(task)
    missing = [path.name for path in image_files if not path.exists()]
    if missing:
        raise ValueError(f"请先生成全部九张图，缺少：{', '.join(missing)}")

    main_file = sku_dir / "主图.png"
    logo_file = sku_dir / "主图_带logo.png"
    print(
        "DEBUG recollage paths:",
        f"sku={task.get('sku')}",
        f"image_folder={sku_dir}",
        f"main_file={main_file}",
        f"logo_file={logo_file}",
        flush=True,
    )
    with connect() as conn:
        conn.execute(
            "UPDATE tasks SET status = 'collaging', generation_progress = '拼图中', updated_at = ? WHERE id = ?",
            (now_iso(), task["id"]),
        )
        log_event(conn, task["id"], task["sku"], "collage_retry_start", "重新执行九宫格拼图")

    compose_nine_grid(image_files, main_file)
    print(f"DEBUG recollage saved: sku={task.get('sku')} file={main_file}", flush=True)
    final_file = main_file
    add_logo = str(task.get("add_logo") or "").strip().lower()
    should_add_logo = add_logo in ("是", "yes", "y", "true", "1")
    if should_add_logo:
        logo_path = task.get("logo_path")
        if not logo_path:
            raise FileNotFoundError("Logo 文件未找到")
        add_logo_to_image(main_file, logo_path, logo_file)
        final_file = logo_file
        print(f"DEBUG recollage logo saved: sku={task.get('sku')} file={logo_file}", flush=True)

    if not final_file.exists():
        raise RuntimeError("九宫格主图未生成")

    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = 'waiting_confirm', image_path = ?, image_folder = ?, error_message = NULL,
                generation_progress = NULL, generated_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (str(final_file), str(sku_dir), now_iso(), now_iso(), task["id"]),
        )
        log_event(conn, task["id"], task["sku"], "collage_retry_success", str(final_file))
    print(
        "DEBUG db image fields updated by recollage:",
        f"sku={task.get('sku')}",
        f"image_path={final_file}",
        f"image_folder={sku_dir}",
        flush=True,
    )
    return final_file


def rerun_collage_for_sku(sku):
    task = get_task_by_sku(str(sku))
    if not task:
        raise ValueError("Task not found")
    try:
        final_file = apply_collage_for_task(task)
        return {"ok": True, "image_path": str(final_file), "image_url": output_url_for_path(final_file)}
    except Exception as exc:
        message = str(exc)
        with connect() as conn:
            conn.execute(
                """
                UPDATE tasks
                SET status = 'collaging_failed', error_message = ?, generation_progress = NULL, updated_at = ?
                WHERE id = ?
                """,
                (message[:4000], now_iso(), task["id"]),
            )
            log_event(conn, task["id"], task["sku"], "collage_retry_error", message[:800])
        raise


<<<<<<< HEAD
DETAIL_PROMPT_TYPES = [
    ("detail_1.png", "平铺展示图 — 产品完整展开，适合详情页首屏"),
    ("detail_2.png", "近景细节图 — 材质、纹理特写"),
    ("detail_3.png", "结构展示图 — 侧面、背面或边缘轮廓，产品不变形"),
    ("detail_4.png", "功能点静物展示 — 口袋、拉链、扣具、弹性或工艺细节"),
    ("detail_5.png", "卖点组合图 — 多角度静物排列，不出现真人"),
    ("detail_6.png", "氛围静物图 — 产品平铺在干净背景或道具场景中"),
]

UPPER_BODY_DETAIL_PROMPT_TYPES = [
    ("detail_1.png", "上身效果图 — 真人模特穿着展示裤子整体版型"),
    ("detail_2.png", "面料质感图 — 穿着状态下展示丝滑垂感与纹理"),
    ("detail_3.png", "版型展示图 — 侧面或背面展示贴合身体和裤型线条"),
    ("detail_4.png", "功能点展示图 — 真人穿着时展示口袋、腰头、拉链或弹性细节"),
    ("detail_5.png", "搭配展示图 — 模特穿着的简洁商业搭配效果"),
    ("detail_6.png", "氛围展示图 — 纯白或干净背景下的真人模特穿着商业摄影"),
]


def detail_mode_for_task(task):
    mode = str(task.get("detail_mode") or task.get("详情页模式") or "").strip()
    return "上身效果" if mode == "上身效果" else "平铺细节"


def split_detail_prompt_segments(text):
    if "|||" not in str(text or ""):
        return []
    return [segment.strip() for segment in str(text).split("|||") if segment.strip()]


def detail_segment_for_index(text, index):
    segments = split_detail_prompt_segments(text)
    if not segments:
        return ""
    if index <= len(segments):
        return segments[index - 1]
    return segments[-1]


def build_detail_prompt(task, detail_kind, index):
    base_prompt = task.get("prompt") or task.get("product_name") or task.get("sku")
    custom_detail_prompt = (task.get("detail_prompt") or "").strip()
    if custom_detail_prompt:
        custom_segment = detail_segment_for_index(custom_detail_prompt, index)
        if custom_segment:
            return custom_segment
        return custom_detail_prompt
    detail_mode = detail_mode_for_task(task)
    scene = detail_segment_for_index(task.get("scene_description"), index) or (task.get("scene_description") or "").strip()
    style = (task.get("detail_style") or "").strip()
    if detail_mode == "上身效果":
        scene_text = scene or "纯白背景或干净棚拍背景，真人模特穿着展示，突出长裤上身版型、贴合身体的线条、面料丝滑垂感"
        style_text = style or "真实商业摄影风格，清晰产品焦点，自然站姿，干净高级"
        return f"""
Use case: ecommerce product detail page model-wearing image
Base product prompt: {base_prompt}
Detail image {index}: {detail_kind}
Scene requirement: {scene_text}
Visual style: {style_text}
Composition: square 1024x1024, full or partial body model wearing the pants, clear product focus, pure white or clean studio background, realistic lighting, no watermark, no irrelevant text.
Positive requirements: model wearing display, on-body effect, real human model wearing the product, show pants fit, smooth silky fabric texture, natural drape and material details.
Keep product identity, colors, materials and brand elements consistent with the white-background images.
""".strip()
    scene_text = scene or "干净电商详情页背景，可使用桌面、布面、浅色道具或纯色背景，突出产品本身"
    style_text = style or "简洁、高级、真实商业静物摄影风格"
    return f"""
Use case: ecommerce product detail page image
Base product prompt: {base_prompt}
Detail image {index}: {detail_kind}
Scene requirement: {scene_text}
Visual style: {style_text}
Composition: square 1024x1024, flat lay or close-up still-life product photography, clear product focus, realistic lighting, no watermark, no irrelevant text.
Texture requirements: close-up fabric texture, smooth material surface, textile detail display, stitching and craft details when applicable.
Keep product identity, colors, materials and brand elements consistent with the white-background images.
Strict constraints: no human, no model, no mannequin, no body parts, no wearing-on-body scene, no outfit lookbook, no lifestyle person shot.
""".strip()


def detail_count_for_task(task):
    return parse_int_field(task.get("detail_count"), default=6, minimum=1, maximum=20)


def detail_prompt_type(index, detail_mode="平铺细节"):
    prompt_types = UPPER_BODY_DETAIL_PROMPT_TYPES if detail_mode == "上身效果" else DETAIL_PROMPT_TYPES
    if 1 <= index <= len(prompt_types):
        return prompt_types[index - 1]
    return (f"detail_{index}.png", f"补充详情图 — 第 {index} 张产品静物细节展示")


def run_detail_for_task(task):
    sku_dir = image_folder_for_task(task)
    sku_dir.mkdir(parents=True, exist_ok=True)
    started = now_iso()
    detail_count = detail_count_for_task(task)
    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = 'detail_generating', detail_status = 'detail_generating',
                detail_error = NULL, detail_progress = ?, updated_at = ?
            WHERE id = ?
            """,
            (f"详情页生成中(0/{detail_count})", started, task["id"]),
        )
        log_event(conn, task["id"], task["sku"], "detail_start", "Detail image generation started")
    try:
        for index in range(1, detail_count + 1):
            filename, detail_kind = detail_prompt_type(index, detail_mode_for_task(task))
            image_file = sku_dir / filename
            with connect() as conn:
                conn.execute(
                    "UPDATE tasks SET detail_progress = ?, updated_at = ? WHERE id = ?",
                    (f"详情页生成中({index}/{detail_count})", now_iso(), task["id"]),
                )
            if image_file.exists():
                print(f"DEBUG detail image exists, skipped: sku={task.get('sku')} file={image_file}", flush=True)
            else:
                prompt = build_detail_prompt(task, detail_kind, index)
                generate_single_image(prompt, image_file, task.get("reference_image_path") or "")
            if not image_file.exists():
                raise RuntimeError(f"第 {index} 张详情页图片未生成")
            print(f"DEBUG detail image saved: sku={task.get('sku')} file={image_file}", flush=True)
        finished = now_iso()
        with connect() as conn:
            conn.execute(
                """
                UPDATE tasks
                SET status = 'detail_done', detail_status = 'detail_done', detail_error = NULL,
                    detail_progress = NULL, detail_generated_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (finished, finished, task["id"]),
            )
            log_event(conn, task["id"], task["sku"], "detail_success", str(sku_dir))
        return True
    except Exception as exc:
        message = openai_error_message(exc)
        with connect() as conn:
            conn.execute(
                """
                UPDATE tasks
                SET status = 'detail_failed', detail_status = 'detail_failed',
                    detail_error = ?, detail_progress = NULL, updated_at = ?
                WHERE id = ?
                """,
                (message[:4000], now_iso(), task["id"]),
            )
            log_event(conn, task["id"], task["sku"], "detail_error", message[:800])
        return False


def tasks_for_detail_generation(mode, skus):
    with connect() as conn:
        if mode == "selected":
            selected = [str(sku) for sku in (skus or []) if str(sku).strip()]
            if not selected:
                return []
            placeholders = ",".join("?" for _ in selected)
            rows = conn.execute(
                f"""
                SELECT * FROM tasks
                WHERE sku IN ({placeholders})
                  AND status IN ('waiting_confirm', 'awaiting_confirm', 'detail_failed', 'detail_done')
                ORDER BY id ASC
                """,
                selected,
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT * FROM tasks
                WHERE status IN ('waiting_confirm', 'awaiting_confirm', 'detail_failed')
                ORDER BY id ASC
                """
            ).fetchall()
    return [dict(row) for row in rows]


def detail_generation_worker(tasks):
    global detail_worker_running
    detail_progress_state.update({"state": "running", "completed": 0, "total": len(tasks), "message": "详情页生成中"})
    try:
        success = 0
        failed = 0
        for task in tasks:
            ok = run_detail_for_task(task)
            if ok:
                success += 1
            else:
                failed += 1
            detail_progress_state.update(
                {
                    "completed": success + failed,
                    "success": success,
                    "failed": failed,
                    "message": f"正在生成详情页：{success + failed}/{len(tasks)}",
                }
            )
            time.sleep(IMAGE_DELAY_SECONDS)
        detail_progress_state.update({"state": "idle", "message": f"详情页完成：成功 {success}，失败 {failed}"})
    finally:
        with detail_lock:
            detail_worker_running = False


def start_detail_generation(mode="all", skus=None):
    global detail_worker_running
    tasks = tasks_for_detail_generation(mode, skus or [])
    if not tasks:
        return {"ok": False, "error": "没有可生成详情页的 SKU"}
    with detail_lock:
        if detail_worker_running:
            return {"ok": False, "error": "详情页正在生成中，请稍后"}
        detail_worker_running = True
    thread = threading.Thread(target=detail_generation_worker, args=(tasks,), daemon=True)
    thread.start()
    return {"ok": True, "total": len(tasks), "success": 0, "failed": 0}


def detail_status():
    status = dict(detail_progress_state)
    status["running"] = detail_worker_running
    return status


=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
def mark_generation_error(task, message):
    current_status = get_task(task["id"])
    is_collage_stage = current_status and current_status.get("status") == "collaging"
    status = "quota_exhausted" if is_quota_error(message) else ("collaging_failed" if is_collage_stage else "failed")
    if status == "quota_exhausted":
        set_quota("exhausted", "OpenAI API 返回额度不足或频率限制")
    print(
        "DEBUG generation failed:",
        f"sku={task.get('sku')}",
        f"status={status}",
        f"reason={message[:800]}",
        flush=True,
    )
    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = ?, error_message = ?, updated_at = ?, retry_count = retry_count + 1
            WHERE id = ?
            """,
            (status, message[:4000], now_iso(), task["id"]),
        )
        log_event(conn, task["id"], task["sku"], "generation_error", f"{status}: {message[:800]}")


def set_quota(status, message):
    quota_state.update({"status": status, "message": message, "updated_at": now_iso()})


def next_pending_task():
    with connect() as conn:
        row = conn.execute(
            """
            SELECT * FROM tasks
            WHERE status = 'pending'
            ORDER BY id ASC
            LIMIT 1
            """
        ).fetchone()
        return dict(row) if row else None


def generation_worker():
    global worker_running
    with worker_lock:
        if worker_running:
            return
        worker_running = True
        queue_control["state"] = "running"
        queue_control["pause_requested"] = False
        queue_control["stop_requested"] = False
    try:
        while True:
            if queue_control["stop_requested"] or queue_control["pause_requested"]:
                return
            task = next_pending_task()
            if not task:
                queue_control["state"] = "idle"
                return
            ok = run_image_for_task(task)
            if not ok:
                latest = get_task(task["id"])
                if latest and latest["status"] == "quota_exhausted":
                    queue_control["state"] = "paused"
                    return
            if queue_control["stop_requested"] or queue_control["pause_requested"]:
                return
            time.sleep(IMAGE_DELAY_SECONDS)
    finally:
        with worker_lock:
            worker_running = False
            if queue_control["stop_requested"]:
                queue_control["state"] = "stopped"
            elif queue_control["pause_requested"]:
                queue_control["state"] = "paused"
            elif queue_control["state"] == "running":
                queue_control["state"] = "idle"


def start_worker():
    queue_control["state"] = "running"
    queue_control["pause_requested"] = False
    queue_control["stop_requested"] = False
    thread = threading.Thread(target=generation_worker, daemon=True)
    thread.start()


def pause_generation():
    if worker_running:
        queue_control["pause_requested"] = True
        queue_control["stop_requested"] = False
        queue_control["state"] = "paused"
    return generate_status()


def stop_generation():
    queue_control["stop_requested"] = True
    queue_control["pause_requested"] = False
    queue_control["state"] = "stopped"
    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = 'pending', updated_at = ?
            WHERE status = 'generating'
            """,
            (now_iso(),),
        )
    return generate_status()


def generate_status():
    with connect() as conn:
        counts = {
            row["status"]: row["count"]
            for row in conn.execute("SELECT status, COUNT(*) AS count FROM tasks GROUP BY status")
        }
    pending = counts.get("pending", 0)
    generating = counts.get("generating", 0)
    total = sum(counts.values())
    completed = max(total - pending - generating, 0)
    state = queue_control["state"]
    if worker_running and not queue_control["pause_requested"] and not queue_control["stop_requested"]:
        state = "running"
    elif not worker_running and state == "running":
        state = "idle"
    return {
        "state": state,
        "running": worker_running,
        "paused": state == "paused",
        "pending": pending,
        "generating": generating,
        "completed": completed,
        "total": total,
    }


def queue_all_paused():
    with connect() as conn:
        conn.execute(
            """
            UPDATE tasks
            SET status = 'pending', error_message = NULL, updated_at = ?
            WHERE status = 'quota_exhausted'
            """,
            (now_iso(),),
        )
        conn.execute(
            """
            INSERT INTO events(task_id, sku, event_type, message, created_at)
            SELECT id, sku, 'resume_quota', 'Quota paused task queued again', ?
            FROM tasks
            WHERE status = 'pending'
            """,
            (now_iso(),),
        )
    set_quota("unknown", "已加入待续跑队列，等待下一次 API 结果")
    start_worker()


def get_task(task_id):
    with connect() as conn:
        row = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        return dict(row) if row else None


def get_task_by_sku(sku):
    with connect() as conn:
        row = conn.execute("SELECT * FROM tasks WHERE sku = ?", (sku,)).fetchone()
        return dict(row) if row else None


def image_path_for_task(task):
    if task and task.get("image_path"):
        return Path(task["image_path"])
    return None


def download_image_path_for_sku(sku):
    task = get_task_by_sku(sku)
    if not task:
        return None
    return main_image_path_for_task(task)


def main_image_path_for_task(task):
    candidates = []
    if task.get("image_path"):
        candidates.append(Path(task["image_path"]))
    if task.get("image_folder"):
        candidates.append(Path(task["image_folder"]) / "主图_带logo.png")
        candidates.append(Path(task["image_folder"]) / "主图.png")
        candidates.append(Path(task["image_folder"]) / "main.png")
    candidates.append(image_folder_for_task(task) / "主图_带logo.png")
    candidates.append(image_folder_for_task(task) / "主图.png")
    candidates.append(image_folder_for_task(task) / "main.png")
    candidates.append(legacy_image_folder_for_task(task) / "主图_带logo.png")
    candidates.append(legacy_image_folder_for_task(task) / "主图.png")
    candidates.append(legacy_image_folder_for_task(task) / "main.png")
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


<<<<<<< HEAD
def image_files_for_task(task, image_type="all"):
=======
def image_files_for_task(task):
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    if not task:
        return []
    folders = []
    folders.append(image_folder_for_task(task))
    if task.get("image_folder"):
        folders.append(Path(task["image_folder"]))
    folders.append(legacy_image_folder_for_task(task))
    files = []
    seen = set()

<<<<<<< HEAD
    if image_type in ("all", "white"):
        for index in range(1, 10):
            for folder in folders:
                candidate = folder / f"{index}.png"
                if candidate.exists():
                    resolved = candidate.resolve()
                    if resolved not in seen:
                        files.append((f"{index}.png", candidate))
                        seen.add(resolved)
                    break

        main_path = main_image_path_for_task(task)
        if main_path and main_path.exists():
            resolved = main_path.resolve()
            if resolved not in seen:
                files.append(("主图.png", main_path))

    if image_type in ("all", "detail"):
        for index in range(1, detail_count_for_task(task) + 1):
            filename = f"detail_{index}.png"
            for folder in folders:
                candidate = folder / filename
                if candidate.exists():
                    resolved = candidate.resolve()
                    if resolved not in seen:
                        files.append((filename, candidate))
                        seen.add(resolved)
                    break
=======
    for index in range(1, 10):
        for folder in folders:
            candidate = folder / f"{index}.png"
            if candidate.exists():
                resolved = candidate.resolve()
                if resolved not in seen:
                    files.append((f"{index}.png", candidate))
                    seen.add(resolved)
                break

    main_path = main_image_path_for_task(task)
    if main_path and main_path.exists():
        resolved = main_path.resolve()
        if resolved not in seen:
            files.append(("主图.png", main_path))
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e

    return files


def folder_with_complete_nine_images(task):
    folders = [image_folder_for_task(task)]
    if task.get("image_folder"):
        folders.append(Path(task["image_folder"]))
    folders.append(legacy_image_folder_for_task(task))
    seen = set()
    for folder in folders:
        resolved = folder.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        image_files = [folder / f"{index}.png" for index in range(1, 10)]
        if all(path.exists() for path in image_files):
            return folder, image_files
    return folders[0], [folders[0] / f"{index}.png" for index in range(1, 10)]


def output_url_for_path(path):
    resolved = Path(path).resolve()
    output_root = OUTPUT_DIR.resolve()
    if resolved == output_root:
        return "/output"
    if output_root not in resolved.parents:
        return ""
    relative = resolved.relative_to(output_root)
    return "/output/" + "/".join(quote(part) for part in relative.parts)


def image_list_for_sku(sku):
    task = get_task_by_sku(str(sku))
<<<<<<< HEAD
    files = image_files_for_task(task, "all")
=======
    files = image_files_for_task(task)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    return [
        {
            "name": name,
            "url": output_url_for_path(path),
            "path": str(path),
        }
        for name, path in files
    ]


<<<<<<< HEAD
def zip_all_images_for_sku(sku, image_type="all"):
    task = get_task_by_sku(str(sku))
    files = image_files_for_task(task, image_type)
=======
def zip_all_images_for_sku(sku):
    task = get_task_by_sku(str(sku))
    files = image_files_for_task(task)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    if not files:
        raise ValueError("该 SKU 暂无可下载图片")

    EXPORT_DIR.mkdir(exist_ok=True)
    zip_path = EXPORT_DIR / zip_filename_for_skus([sku])
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for arcname, path in files:
            archive.write(path, arcname=arcname)
    return zip_path


<<<<<<< HEAD
def zip_images_for_skus(skus, image_type="all"):
=======
def zip_images_for_skus(skus):
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    EXPORT_DIR.mkdir(exist_ok=True)
    filename = zip_filename_for_skus(skus)
    zip_path = EXPORT_DIR / filename
    added = 0
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for sku in skus:
            task = get_task_by_sku(str(sku))
            folder_name = download_base_name_for_task(task, str(sku))
<<<<<<< HEAD
            for arcname, image_path in image_files_for_task(task, image_type):
=======
            for arcname, image_path in image_files_for_task(task):
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                archive.write(image_path, arcname=f"{folder_name}/{arcname}")
                added += 1
    if added == 0:
        raise ValueError("选中的 SKU 暂无可下载图片")
    return zip_path


def delete_tasks_by_skus(skus):
    deleted = 0
    with connect() as conn:
        for sku in skus:
            task = conn.execute("SELECT * FROM tasks WHERE sku = ?", (str(sku),)).fetchone()
            if not task:
                continue
            image_folder = task["image_folder"]
            image_path = task["image_path"]
            conn.execute("DELETE FROM events WHERE sku = ?", (task["sku"],))
            conn.execute("DELETE FROM tasks WHERE sku = ?", (task["sku"],))
            deleted += 1
            folder = Path(image_folder) if image_folder else (Path(image_path).parent if image_path else image_folder_for_task(dict(task)))
            if folder and folder.exists() and OUTPUT_DIR.resolve() in folder.resolve().parents:
                shutil.rmtree(folder, ignore_errors=True)
    return {"ok": True, "deleted": deleted}


def export_generation_status(status):
    if status in ("waiting_confirm", "awaiting_confirm", "approved"):
        return "成功"
<<<<<<< HEAD
    if status == "detail_done":
        return "已完成"
    if status == "detail_generating":
        return "生成中"
    if status in ("failed", "collaging_failed", "quality_failed", "rejected", "quota_exhausted"):
        return "失败"
    if status == "detail_failed":
        return "失败"
    return "待生成"


def task_output_path_text(row):
    task = dict(row)
    main_path = main_image_path_for_task(task)
    if main_path:
        return str(main_path)
    return row["image_path"] or row["image_folder"] or ""


def task_export_log(row):
    data = dict(row)
    return data.get("detail_error") or data.get("error_message") or ""


def find_ws_header_row(ws, required_headers):
    required = set(required_headers)
    for row_index in range(1, ws.max_row + 1):
        values = {str(ws.cell(row_index, col).value or "").strip() for col in range(1, ws.max_column + 1)}
        if required & values:
            return row_index
    return 1


def header_map_for_ws(ws, header_row):
    headers = {}
    for col in range(1, ws.max_column + 1):
        header = normalize_excel_column_name(ws.cell(header_row, col).value)
        if header and header not in headers:
            headers[header] = col
    return headers


def detect_header_row_for_columns(ws, columns, max_scan_rows=10):
    target_headers = {normalize_excel_column_name(column) for column in columns}
    best_row = 1
    best_score = 0
    max_row = min(ws.max_row, max_scan_rows)
    for row_index in range(1, max_row + 1):
        row_headers = {
            normalize_excel_column_name(ws.cell(row_index, col).value)
            for col in range(1, ws.max_column + 1)
        }
        score = len(target_headers & row_headers)
        if score > best_score:
            best_score = score
            best_row = row_index
    return best_row


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.font:
        target_cell.font = copy(source_cell.font)
    if source_cell.fill:
        target_cell.fill = copy(source_cell.fill)
    if source_cell.border:
        target_cell.border = copy(source_cell.border)
    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)


def ensure_columns(ws, header_row, columns):
    if header_row is None:
        header_row = detect_header_row_for_columns(ws, columns)
    headers = header_map_for_ws(ws, header_row)
    for column_name in columns:
        normalized_column_name = normalize_excel_column_name(column_name)
        if normalized_column_name in headers:
            headers[column_name] = headers[normalized_column_name]
            continue
        new_col = max(headers.values(), default=0) + 1
        header_cell = ws.cell(header_row, new_col, column_name)
        if new_col > 1:
            copy_cell_style(ws.cell(header_row, new_col - 1), header_cell)
            previous_column_letter = get_column_letter(new_col - 1)
            new_column_letter = get_column_letter(new_col)
            width = ws.column_dimensions[previous_column_letter].width
            if width:
                ws.column_dimensions[new_column_letter].width = width
        headers[column_name] = new_col
    return headers


def write_backfill_row(ws, row_index, headers, task, field_map):
    for header, value in field_map.items():
        col = headers[header]
        cell = ws.cell(row_index, col)
        if row_index > 1 and not cell.has_style:
            copy_cell_style(ws.cell(row_index - 1, col), cell)
        cell.value = value


def task_backfill_values(row):
    return {
        "SKU": row["sku"] or "",
        "是否添加品牌logo": row["add_logo"] or "",
        "logo文件地址": row["logo_path"] or "",
        "需校验品牌文字/元素": row["brand_check"] or "",
        "优化后提示词": row["optimized_prompt"] or "",
        "AI推荐标题": row["ai_title"] or "",
        "生成时间": row["detail_generated_at"] or row["generated_at"] or "",
        "输出文件地址": task_output_path_text(row),
        "生成图片张数": len(image_files_for_task(dict(row), "all")),
        "图片颜色列表": row["color_name"] or "",
        "执行状态": export_generation_status(row["status"]),
        "执行日志": task_export_log(row),
    }


def old_task_backfill_values(row):
    return {
        "优化后提示词": row["optimized_prompt"] or "",
        "提示词原文快照": row["prompt_snapshot"] or "",
        "生成状态": export_generation_status(row["status"]),
        "失败原因": task_export_log(row),
        "生成时间": row["detail_generated_at"] or row["generated_at"] or "",
        "输出文件地址": task_output_path_text(row),
    }


def build_task_lookup(rows):
    lookup = {}
    for row in rows:
        for key in (row["sku"], row["spu"]):
            key = str(key or "").strip()
            if key and key not in lookup:
                lookup[key] = row
    return lookup


def export_source_status():
    source_path_text = get_config("last_import_path")
    exists = bool(source_path_text and Path(source_path_text).exists())
    return {"exists": exists, "path": source_path_text or ""}


IMAGE_GENERATION_SHEET = "图片生成"
IMAGE_GENERATION_COLUMNS = [
    "SKU",
    "是否添加品牌logo",
    "logo文件地址",
    "需校验品牌文字/元素",
    "优化后提示词",
    "AI推荐标题",
    "生成时间",
    "输出文件地址",
    "生成图片张数",
    "图片颜色列表",
    "执行状态",
    "执行日志",
]


def first_row_headers(ws):
    return [str(ws.cell(1, col).value or "").strip() for col in range(1, ws.max_column + 1)]


def recreate_image_generation_sheet(wb):
    if IMAGE_GENERATION_SHEET in wb.sheetnames:
        ws_old = wb[IMAGE_GENERATION_SHEET]
        index = wb.sheetnames.index(IMAGE_GENERATION_SHEET)
        wb.remove(ws_old)
        ws = wb.create_sheet(IMAGE_GENERATION_SHEET, index)
    else:
        ws = wb.create_sheet(IMAGE_GENERATION_SHEET)
    for col, header in enumerate(IMAGE_GENERATION_COLUMNS, start=1):
        ws.cell(1, col, header)
    return ws


def print_export_debug(wb, ws_board, headers, rows, excel_skus, matched_skus):
    db_skus = [str(row["sku"] or "").strip() for row in rows if str(row["sku"] or "").strip()]
    missing_in_excel = [sku for sku in db_skus if sku not in excel_skus]
    extra_in_excel = [sku for sku in excel_skus if sku and sku not in set(db_skus)]
    print(f"导出调试 Sheet 列表: {wb.sheetnames}", flush=True)
    print(f"导出调试 是否存在「{IMAGE_GENERATION_SHEET}」Sheet: {IMAGE_GENERATION_SHEET in wb.sheetnames}", flush=True)
    if ws_board:
        print(f"导出调试「{IMAGE_GENERATION_SHEET}」第一行列头: {first_row_headers(ws_board)}", flush=True)
    print(f"导出调试「{IMAGE_GENERATION_SHEET}」识别到的表头: {list(headers.keys()) if headers else []}", flush=True)
    print(f"导出调试 数据库 SKU 数量: {len(db_skus)}，示例: {db_skus[:20]}", flush=True)
    print(f"导出调试 Excel SKU 数量: {len(excel_skus)}，示例: {excel_skus[:20]}", flush=True)
    print(f"导出调试 匹配 SKU 数量: {len(matched_skus)}，示例: {sorted(matched_skus)[:20]}", flush=True)
    print(f"导出调试 数据库有但 Excel 没有的 SKU 数量: {len(missing_in_excel)}，示例: {missing_in_excel[:20]}", flush=True)
    print(f"导出调试 Excel 有但数据库没有的 SKU 数量: {len(extra_in_excel)}，示例: {extra_in_excel[:20]}", flush=True)


def export_excel_from_source(rows, table_mode, output_path, source_path_override=None):
    source_path_text = str(source_path_override) if source_path_override else get_config("last_import_path")
    if not source_path_text:
        return False
    source_path = Path(source_path_text)
    if not source_path.exists():
        print(f"原始 Excel 未找到：{source_path}", flush=True)
        return False

    wb = load_workbook(source_path)
    resolved_mode = get_config("last_import_table_mode") or table_mode
    task_lookup = build_task_lookup(rows)

    if resolved_mode == "old":
        ws, _ = select_import_sheet(wb, "old")
        header_row = find_ws_header_row(ws, ["SPU", "SKU", "图片生成提示词"])
        headers = ensure_columns(
            ws,
            header_row,
            [
                "SPU", "图片生成提示词", "图片颜色", "图片命名", "是否添加品牌logo", "logo文件地址",
                "需校验品牌文字/元素", "参考图地址", "场景描述", "详情页提示词", "详情页风格", "详情页数量",
                "优化后提示词", "提示词原文快照", "生成状态", "失败原因", "生成时间", "输出文件地址",
            ],
        )
        matched = set()
        for row_index in range(header_row + 1, ws.max_row + 1):
            key_values = []
            for header in ("SKU", "SPU"):
                if header in headers:
                    key_values.append(str(ws.cell(row_index, headers[header]).value or "").strip())
            task = next((task_lookup.get(key) for key in key_values if key), None)
            if not task:
                continue
            write_backfill_row(ws, row_index, headers, task, old_task_backfill_values(task))
            matched.add(task["sku"])
        for task in rows:
            if task["sku"] in matched:
                continue
            row_index = ws.max_row + 1
            if "SKU" in headers:
                ws.cell(row_index, headers["SKU"], task["sku"])
            if "SPU" in headers:
                ws.cell(row_index, headers["SPU"], task["spu"] or task["sku"] or "")
            write_backfill_row(ws, row_index, headers, task, old_task_backfill_values(task))
        wb.save(output_path)
        return True

    print(f"导出调试 原始 Excel Sheet 列表: {wb.sheetnames}", flush=True)
    print(f"导出调试 是否存在「{IMAGE_GENERATION_SHEET}」Sheet: {IMAGE_GENERATION_SHEET in wb.sheetnames}", flush=True)
    if IMAGE_GENERATION_SHEET in wb.sheetnames:
        ws_board = wb[IMAGE_GENERATION_SHEET]
        print(f"导出调试「{IMAGE_GENERATION_SHEET}」第一行列头: {first_row_headers(ws_board)}", flush=True)
    else:
        print(f"导出调试 原始 Excel 中没有「{IMAGE_GENERATION_SHEET}」Sheet，将新建该 Sheet", flush=True)
        ws_board = recreate_image_generation_sheet(wb)

    header_row = find_ws_header_row(ws_board, ["SKU"])
    initial_headers = header_map_for_ws(ws_board, header_row)
    if "SKU" not in initial_headers:
        print(f"导出调试「{IMAGE_GENERATION_SHEET}」Sheet 列名不匹配：找不到 SKU 列，将重建该 Sheet", flush=True)
        ws_board = recreate_image_generation_sheet(wb)
        header_row = 1
    headers = ensure_columns(
        ws_board,
        header_row,
        IMAGE_GENERATION_COLUMNS,
    )
    sku_col = headers["SKU"]
    excel_skus = []
    matched_skus = set()
    for row_index in range(header_row + 1, ws_board.max_row + 1):
        sku = str(ws_board.cell(row_index, sku_col).value or "").strip()
        if sku:
            excel_skus.append(sku)
        task = task_lookup.get(sku)
        if not task:
            continue
        write_backfill_row(ws_board, row_index, headers, task, task_backfill_values(task))
        matched_skus.add(task["sku"])

    original_excel_skus = list(excel_skus)
    original_matched_skus = set(matched_skus)
    print_export_debug(wb, ws_board, headers, rows, original_excel_skus, original_matched_skus)

    appended = 0
    for task in rows:
        if task["sku"] in matched_skus:
            continue
        row_index = ws_board.max_row + 1
        write_backfill_row(ws_board, row_index, headers, task, task_backfill_values(task))
        excel_skus.append(task["sku"])
        matched_skus.add(task["sku"])
        appended += 1

    print(f"导出调试 追加未匹配 SKU 行数: {appended}", flush=True)

    if "商品总表" in wb.sheetnames:
        ws_total = wb["商品总表"]
    else:
        ws_total = wb[wb.sheetnames[0]]
    total_header_row = find_ws_header_row(ws_total, ["SKU", "图片生成提示词"])
    ensure_columns(
        ws_total,
        total_header_row,
        [
            "SKU", "SPU", "图片生成提示词", "图片颜色", "图片命名", "是否添加品牌logo", "logo文件地址",
            "需校验品牌文字/元素", "参考图地址", "场景描述", "详情页提示词", "详情页风格", "详情页数量",
        ],
    )
    wb.save(output_path)
    return True


def generated_export_excel(rows, table_mode, path):
    wb = Workbook()
    if table_mode == "old":
        ws = wb.active
        ws.title = "任务表"
        columns = [
            "SPU", "图片生成提示词", "图片颜色", "图片命名", "是否添加品牌logo", "logo文件地址",
            "需校验品牌文字/元素", "参考图地址", "场景描述", "详情页提示词", "详情页风格", "详情页数量", "优化后提示词",
            "提示词原文快照", "生成状态", "失败原因", "生成时间", "输出文件地址", "SKU",
        ]
        ws.append(columns)
        for row in rows:
            ws.append(
                [
                    row["spu"] or row["sku"] or "",
                    row["prompt"] or "",
                    row["color_name"] or "",
                    row["image_name"] or "",
                    row["add_logo"] or "",
                    row["logo_path"] or "",
                    row["brand_check"] or "",
                    row["reference_image_path"] or "",
                    row["scene_description"] or "",
                    row["detail_prompt"] or "",
                    row["detail_style"] or "",
                    row["detail_count"] or 6,
                    row["optimized_prompt"] or "",
                    row["prompt_snapshot"] or "",
                    export_generation_status(row["status"]),
                    task_export_log(row),
                    row["detail_generated_at"] or row["generated_at"] or "",
                    task_output_path_text(row),
                    row["sku"],
                ]
            )
        wb.save(path)
        return path

    ws_total = wb.active
    ws_total.title = "商品总表"
    ws_total.append(["SKU", "SPU", "商品名称", "图片颜色", "图片生成提示词", "参考图地址", "场景描述", "详情页提示词", "详情页风格", "详情页数量"])
    ws_board = wb.create_sheet(IMAGE_GENERATION_SHEET)
    ws_board.append(IMAGE_GENERATION_COLUMNS)
    for row in rows:
        ws_total.append([
            row["sku"],
            row["spu"] or row["sku"] or "",
            row["product_name"] or "",
            row["color_name"] or "",
            row["prompt"] or "",
            row["reference_image_path"] or "",
            row["scene_description"] or "",
            row["detail_prompt"] or "",
            row["detail_style"] or "",
            row["detail_count"] or 6,
        ])
        values = task_backfill_values(row)
        ws_board.append([values.get(column, "") for column in IMAGE_GENERATION_COLUMNS])
    wb.save(path)
    return path


def export_excel(table_mode="new", source_path=None, allow_source_fallback=True):
    EXPORT_DIR.mkdir(exist_ok=True)
    filename = f"erp_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = EXPORT_DIR / filename
    with connect() as conn:
        rows = conn.execute("SELECT * FROM tasks ORDER BY id ASC").fetchall()
    if export_excel_from_source(rows, table_mode, path, source_path):
        return path
    if not allow_source_fallback:
        raise FileNotFoundError("找不到原始 Excel 文件，请手动选择")
    generated_export_excel(rows, table_mode, path)
    return path


excel_lock = threading.RLock()

USER_EXCEL_COLUMNS = [
    "SPU", "平台", "产品名称", "类目", "价格", "币种", "库存", "上架标题", "主题标签",
    "风格参考", "图片颜色", "图片生成提示词", "参考图地址", "场景描述", "详情页模式", "差评词", "店铺",
]
PROGRAM_EXCEL_COLUMNS = [
    "SKU", "是否添加品牌logo", "logo文件地址", "需校验品牌文字/元素", "优化后提示词",
    "AI推荐标题", "生成时间", "输出文件地址", "生成图片张数", "图片颜色列表",
    "执行状态", "执行日志", "是否改色",
]
IMAGE_GENERATION_COLUMNS = PROGRAM_EXCEL_COLUMNS

EXCEL_TO_STATUS = {
    "": "pending",
    "待生成": "pending",
    "待处理": "pending",
    "生成中": "generating",
    "拼图中": "collaging",
    "待确认": "waiting_confirm",
    "已完成": "waiting_confirm",
    "成功": "waiting_confirm",
    "失败": "failed",
    "已通过": "approved",
    "已驳回": "rejected",
    "额度耗尽": "quota_exhausted",
}
STATUS_TO_EXCEL = {
    "pending": "待生成",
    "generating": "生成中",
    "collaging": "拼图中",
    "waiting_confirm": "已完成",
    "awaiting_confirm": "已完成",
    "detail_done": "已完成",
    "approved": "已通过",
    "rejected": "已驳回",
    "failed": "失败",
    "collaging_failed": "失败",
    "quality_failed": "失败",
    "quota_exhausted": "额度耗尽",
}


def active_excel_path():
    path_text = get_config("last_import_path")
    if not path_text:
        return None
    path = Path(path_text)
    return path if path.exists() else None


def backup_excel_file(path):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = path.with_name(f"{path.stem}_备份_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    print(f"Excel 备份完成: {backup_path}", flush=True)
    return backup_path


def excel_status_to_internal(value):
    return EXCEL_TO_STATUS.get(str(value or "").strip(), "pending")


def internal_status_to_excel(value):
    return STATUS_TO_EXCEL.get(str(value or "").strip(), "待生成")


def ensure_excel_columns(ws, columns, header_row=None):
    return ensure_columns(ws, header_row, columns)


def excel_row_dict(ws, headers, row_index):
    return {header: ws.cell(row_index, col).value for header, col in headers.items()}


def find_or_create_board_row(ws, headers, sku, header_row=1):
    sku_col = headers["SKU"]
    for row_index in range(header_row + 1, ws.max_row + 1):
        if str(ws.cell(row_index, sku_col).value or "").strip() == str(sku).strip():
            return row_index
    row_index = ws.max_row + 1
    ws.cell(row_index, sku_col, sku)
    return row_index


def color_list_for_task(task):
    return split_list_text(task.get("color_name") or task.get("图片颜色") or "")


def normalize_recolor_flag(value):
    text = str(value or "").strip().lower()
    if text in ("否", "no", "n", "false", "0"):
        return "否"
    return "是"


def excel_task_from_row(row_index, data):
    sku = str(data.get("SKU") or data.get("SPU") or "").strip()
    spu = str(data.get("SPU") or sku).strip()
    status = excel_status_to_internal(data.get("执行状态"))
    image_folder = str(data.get("输出文件地址") or "").strip()
    if image_folder and Path(image_folder).suffix:
        image_path = image_folder
        image_folder = str(Path(image_folder).parent)
    else:
        image_path = str(Path(image_folder) / "主图.png") if image_folder else ""
    return {
        "id": row_index,
        "sku": sku,
        "spu": spu,
        "product_name": str(data.get("产品名称") or data.get("商品名称") or "").strip(),
        "prompt": str(data.get("图片生成提示词") or "").strip(),
        "color_name": str(data.get("图片颜色列表") or data.get("图片颜色") or "").strip(),
        "image_name": "",
        "add_logo": str(data.get("是否添加品牌logo") or "").strip(),
        "logo_path": str(data.get("logo文件地址") or "").strip(),
        "brand_check": str(data.get("需校验品牌文字/元素") or "").strip(),
        "reference_image_path": str(data.get("参考图地址") or "").strip(),
        "scene_description": str(data.get("场景描述") or "").strip(),
        "detail_mode": detail_mode_for_task(data),
        "detail_style": str(data.get("风格参考") or data.get("详情页风格") or "").strip(),
        "detail_prompt": str(data.get("详情页提示词") or "").strip(),
        "detail_count": parse_int_field(data.get("详情页数量"), default=6, minimum=1, maximum=20),
        "recolor": normalize_recolor_flag(data.get("是否改色")),
        "platform": str(data.get("平台") or "").strip(),
        "style_reference": str(data.get("风格参考") or "").strip(),
        "optimized_prompt": str(data.get("优化后提示词") or "").strip(),
        "ai_title": str(data.get("AI推荐标题") or "").strip(),
        "generated_at": str(data.get("生成时间") or "").strip(),
        "image_folder": image_folder,
        "image_path": image_path,
        "status": status,
        "error_message": str(data.get("执行日志") or "").strip(),
        "generation_progress": None,
        "raw": data,
        "raw_json": json.dumps(data, ensure_ascii=False, default=str),
    }


def excel_load_tasks():
    path = active_excel_path()
    if not path:
        return []
    with excel_lock:
        wb = load_workbook(path)
        if IMAGE_GENERATION_SHEET not in wb.sheetnames:
            return []
        ws = wb[IMAGE_GENERATION_SHEET]
        header_row = detect_header_row_for_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
        headers = ensure_excel_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS, header_row)
        tasks = []
        for row_index in range(header_row + 1, ws.max_row + 1):
            data = excel_row_dict(ws, headers, row_index)
            if not str(data.get("SKU") or data.get("SPU") or "").strip():
                continue
            tasks.append(excel_task_from_row(row_index, data))
        return tasks


def save_task_fields(task, **fields):
    path = active_excel_path()
    if not path:
        raise ValueError("请先导入 Excel")
    with excel_lock:
        wb = load_workbook(path)
        ws = wb[IMAGE_GENERATION_SHEET]
        header_row = detect_header_row_for_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
        headers = ensure_excel_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS, header_row)
        row_index = find_or_create_board_row(ws, headers, task["sku"], header_row)
        for key, value in fields.items():
            if key not in headers:
                headers = ensure_excel_columns(ws, [key], header_row)
            ws.cell(row_index, headers[key], value)
        wb.save(path)


def list_tasks(status=None, q=None):
    tasks = excel_load_tasks()
    if status and status != "all":
        tasks = [task for task in tasks if task["status"] == status]
    if q:
        query = str(q).lower()
        tasks = [
            task for task in tasks
            if query in str(task.get("sku", "")).lower() or query in str(task.get("product_name", "")).lower()
        ]
    for task in tasks:
        task["status_label"] = STATUS_LABELS.get(task["status"], task["status"])
    return sorted(tasks, key=lambda item: item["id"], reverse=True)


def get_task(task_id):
    for task in excel_load_tasks():
        if int(task["id"]) == int(task_id):
            return task
    return None


def get_task_by_sku(sku):
    sku = str(sku)
    for task in excel_load_tasks():
        if str(task["sku"]) == sku:
            return task
    return None


def update_status(task_id, status, message=""):
    task = get_task(task_id)
    if not task:
        raise ValueError("Task not found")
    save_task_fields(task, 执行状态=internal_status_to_excel(status), 执行日志=message or "")


def reset_for_generation(task_id):
    task = get_task(task_id)
    if not task:
        raise ValueError("Task not found")
    save_task_fields(task, 执行状态="待生成", 执行日志="")


def next_pending_task():
    for task in sorted(excel_load_tasks(), key=lambda item: item["id"]):
        if task["status"] == "pending":
            return task
    return None


def excel_init_from_source(path, table_mode="auto"):
    wb = load_workbook(path)
    ws_source, resolved_mode = select_import_sheet(wb, table_mode)
    source_rows = list(ws_source.iter_rows(values_only=True))
    if not source_rows:
        wb.save(path)
        return {"imported": 0, "updated": 0, "skipped": 0, "mode": resolved_mode, "sheet": ws_source.title}
    header_index = find_header_row(source_rows, ["SKU", "SPU", "图片生成提示词"])
    source_headers = ["" if cell is None else str(cell).strip() for cell in source_rows[header_index]]

    if IMAGE_GENERATION_SHEET in wb.sheetnames:
        ws_board = wb[IMAGE_GENERATION_SHEET]
    else:
        ws_board = wb.create_sheet(IMAGE_GENERATION_SHEET)
    board_header_row = detect_header_row_for_columns(ws_board, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
    board_headers = ensure_excel_columns(ws_board, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS, board_header_row)
    existing = {}
    for row_index in range(board_header_row + 1, ws_board.max_row + 1):
        sku = str(ws_board.cell(row_index, board_headers["SKU"]).value or "").strip()
        if sku:
            existing[sku] = row_index

    imported = updated = skipped = 0
    for values in source_rows[header_index + 1:]:
        data = normalize_row(source_headers, list(values) + [None] * (len(source_headers) - len(values)))
        sku = pick_field(data, ["SKU", "SPU", "sku", "spu", "货号", "产品系列编号"])
        if not sku:
            skipped += 1
            continue
        is_existing = sku in existing
        row_index = existing.get(sku) or ws_board.max_row + 1
        existing[sku] = row_index
        for column in USER_EXCEL_COLUMNS:
            value = pick_field(data, [column])
            if value:
                ws_board.cell(row_index, board_headers[column], value)
        ws_board.cell(row_index, board_headers["SKU"], sku)
        if not ws_board.cell(row_index, board_headers["SPU"]).value:
            ws_board.cell(row_index, board_headers["SPU"], sku)
        if not ws_board.cell(row_index, board_headers["执行状态"]).value:
            ws_board.cell(row_index, board_headers["执行状态"], "待生成")
        if not ws_board.cell(row_index, board_headers["是否改色"]).value:
            ws_board.cell(row_index, board_headers["是否改色"], "是")
        if not ws_board.cell(row_index, board_headers["图片颜色列表"]).value:
            ws_board.cell(row_index, board_headers["图片颜色列表"], pick_field(data, ["图片颜色"]))
        if is_existing:
            updated += 1
        else:
            imported += 1
    wb.save(path)
    return {"imported": imported, "updated": updated, "skipped": skipped, "mode": resolved_mode, "sheet": IMAGE_GENERATION_SHEET}


def import_excel(path, table_mode="auto"):
    backup_excel_file(path)
    set_config("last_import_path", str(Path(path).resolve()))
    set_config("last_import_table_mode", table_mode)
    return excel_init_from_source(path, table_mode)


COLOR_MAP = {
    "红": (205, 45, 45), "红色": (205, 45, 45),
    "黑": (35, 35, 35), "黑色": (35, 35, 35),
    "白": (238, 238, 230), "白色": (238, 238, 230),
    "黄": (225, 185, 40), "黄色": (225, 185, 40),
    "蓝": (55, 100, 190), "蓝色": (55, 100, 190),
    "绿": (55, 150, 85), "绿色": (55, 150, 85),
    "紫": (125, 80, 170), "紫色": (125, 80, 170),
    "灰": (125, 125, 125), "灰色": (125, 125, 125),
    "橙": (220, 120, 35), "橙色": (220, 120, 35),
}


COLOR_HUE_MAP = {
    "红": 0, "红色": 0,
    "橙": 20, "橙色": 20,
    "黄": 45, "黄色": 45,
    "绿": 120, "绿色": 120,
    "蓝": 180, "蓝色": 180,
    "紫": 270, "紫色": 270,
}


def color_rgb(color_name):
    text = str(color_name or "").strip()
    return COLOR_MAP.get(text) or next((rgb for key, rgb in COLOR_MAP.items() if key in text), (120, 120, 120))


def color_hue_degrees(color_name):
    text = str(color_name or "").strip()
    for key, hue in COLOR_HUE_MAP.items():
        if key in text:
            return hue
    r, g, b = color_rgb(text)
    hue, _, _ = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
    return int(hue * 360)


def color_family(color_name):
    text = str(color_name or "").strip()
    if "黑" in text:
        return "black"
    if "白" in text:
        return "white"
    if "灰" in text:
        return "gray"
    return "color"


def recolored_hsv_pixel(h, s, v, target_color):
    family = color_family(target_color)
    if family == "black":
        return h, min(s, 45), max(18, min(58, int(v * 0.22)))
    if family == "white":
        return h, 0, max(228, min(255, int(235 + (v - 128) * 0.18)))
    if family == "gray":
        return h, 0, max(85, min(175, int(v * 0.62)))

    target_hue = int(color_hue_degrees(target_color) / 360 * 255) % 256
    new_s = max(70, min(230, int(s * 0.9 + 55)))
    new_v = max(35, min(255, int(v * 0.98 + 4)))
    return target_hue, new_s, new_v


def is_white_background_pixel(s, v):
    return s <= 13 and v >= 242


def is_product_candidate_pixel(s, v, alpha):
    if alpha == 0:
        return False
    if is_white_background_pixel(s, v):
        return False
    return s > 18 or v < 238


def recolor_image(source_file, target_file, color_name, sku="", original_color="红"):
    base = Image.open(source_file).convert("RGBA")
    alpha = base.getchannel("A")
    hsv = base.convert("RGB").convert("HSV")
    pixels = hsv.load()
    alpha_pixels = alpha.load()
    width, height = hsv.size
    affected = 0

    for y in range(height):
        for x in range(width):
            h, s, v = pixels[x, y]
            if not is_product_candidate_pixel(s, v, alpha_pixels[x, y]):
                continue
            pixels[x, y] = recolored_hsv_pixel(h, s, v, color_name)
            affected += 1

    recolored = hsv.convert("RGB").convert("RGBA")
    recolored.putalpha(alpha)
    background = base.copy()
    background.alpha_composite(recolored)
    background.convert("RGB").save(target_file)
    print(
        "DEBUG color replacement:",
        f"sku={sku}",
        f"original_color={original_color}",
        f"target_color={color_name}",
        f"pixels_affected={affected}",
        flush=True,
    )


def title_language_for_platform(platform):
    text = str(platform or "").lower()
    if "ozon" in text:
        return "俄文"
    if "temu" in text:
        return "英文"
    return "中文"


def generate_ai_title(task):
    try:
        language = title_language_for_platform(task.get("platform"))
        prompt = (
            f"请为跨境电商商品生成一个{language}上架标题，只输出标题，不要解释。\n"
            f"产品名称：{task.get('product_name')}\n"
            f"生图提示词：{str(task.get('prompt') or '')[:200]}\n"
            f"颜色列表：{task.get('color_name')}\n"
            f"风格参考：{task.get('style_reference')}\n"
        )
        codex_path = os.path.expandvars(r"%APPDATA%\npm\codex.cmd")
        cmd = [codex_path if Path(codex_path).exists() else "codex", "exec", prompt]
        result = subprocess.run(
            subprocess.list2cmdline(cmd),
            shell=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
        )
        if result.returncode != 0:
            print(f"AI 标题生成失败: {result.stderr or result.stdout}", flush=True)
            return ""
        return (result.stdout or "").strip().splitlines()[-1][:180]
    except Exception as exc:
        print(f"AI 标题生成失败: {exc}", flush=True)
        return ""


def add_manual_task(payload):
    path = active_excel_path()
    if not path:
        raise ValueError("请先导入 Excel，再手动新增 SKU")
    sku = (payload.get("sku") or payload.get("SKU") or payload.get("spu") or payload.get("SPU") or "").strip()
    product_name = (payload.get("product_name") or payload.get("商品名称") or "").strip()
    prompt = (payload.get("prompt") or payload.get("图片生成提示词") or "").strip()
    if not sku or not product_name or not prompt:
        raise ValueError("SKU、商品名称、生图提示词均为必填")
    with excel_lock:
        wb = load_workbook(path)
        ws = wb[IMAGE_GENERATION_SHEET] if IMAGE_GENERATION_SHEET in wb.sheetnames else wb.create_sheet(IMAGE_GENERATION_SHEET)
        header_row = detect_header_row_for_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
        headers = ensure_excel_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS, header_row)
        row_index = find_or_create_board_row(ws, headers, sku, header_row)
        values = {
            "SKU": sku,
            "SPU": payload.get("spu") or payload.get("SPU") or sku,
            "产品名称": product_name,
            "图片生成提示词": prompt,
            "图片颜色": payload.get("图片颜色") or payload.get("color_name") or "",
            "图片颜色列表": payload.get("图片颜色") or payload.get("color_name") or "",
            "执行状态": "待生成",
            "是否改色": normalize_recolor_flag(payload.get("是否改色")),
        }
        for column, value in values.items():
            ws.cell(row_index, headers[column], value)
        wb.save(path)


def delete_tasks_by_skus(skus):
    path = active_excel_path()
    if not path:
        return {"ok": True, "deleted": 0}
    selected = {str(sku).strip() for sku in skus if str(sku).strip()}
    deleted = 0
    with excel_lock:
        wb = load_workbook(path)
        if IMAGE_GENERATION_SHEET not in wb.sheetnames:
            return {"ok": True, "deleted": 0}
        ws = wb[IMAGE_GENERATION_SHEET]
        header_row = detect_header_row_for_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
        headers = ensure_excel_columns(ws, USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS, header_row)
        sku_col = headers["SKU"]
        for row_index in range(ws.max_row, header_row, -1):
            sku = str(ws.cell(row_index, sku_col).value or "").strip()
            if sku in selected:
                task = excel_task_from_row(row_index, excel_row_dict(ws, headers, row_index))
                folder = image_folder_for_task(task)
                if folder.exists() and OUTPUT_DIR.resolve() in folder.resolve().parents:
                    shutil.rmtree(folder, ignore_errors=True)
                ws.delete_rows(row_index, 1)
                deleted += 1
        wb.save(path)
    return {"ok": True, "deleted": deleted}


def tasks_for_detail_generation(mode, skus):
    tasks = excel_load_tasks()
    if mode == "selected":
        selected = {str(sku) for sku in (skus or [])}
        tasks = [task for task in tasks if task["sku"] in selected]
    return [
        task for task in tasks
        if task["status"] in ("waiting_confirm", "awaiting_confirm", "detail_failed", "detail_done")
    ]


def run_detail_for_task(task):
    sku_dir = image_folder_for_task(task)
    sku_dir.mkdir(parents=True, exist_ok=True)
    detail_count = detail_count_for_task(task)
    save_task_fields(task, 执行状态="生成中", 执行日志="")
    try:
        for index in range(1, detail_count + 1):
            filename, detail_kind = detail_prompt_type(index, detail_mode_for_task(task))
            image_file = sku_dir / filename
            if not image_file.exists():
                prompt = build_detail_prompt(task, detail_kind, index)
                generate_single_image(prompt, image_file, task.get("reference_image_path") or "")
        save_task_fields(task, 执行状态="已完成", 生成时间=now_iso(), 执行日志="")
        return True
    except Exception as exc:
        save_task_fields(task, 执行状态="失败", 执行日志=openai_error_message(exc)[:4000])
        return False


def run_image_for_task(task):
    sku_dir = image_folder_for_task(task)
    sku_dir.mkdir(parents=True, exist_ok=True)
    main_file = sku_dir / "主图.png"
    logo_file = sku_dir / "主图_带logo.png"
    optimized_prompt = build_codex_prompt(task, main_file)
    colors = color_list_for_task(task) or ["默认"]
    use_recolor = normalize_recolor_flag(task.get("recolor")) == "是"
    save_task_fields(task, 执行状态="生成中", 执行日志="", 优化后提示词=optimized_prompt, 输出文件地址=str(main_file), 图片颜色列表="/".join(colors))
    try:
        image_files = []
        base_file = sku_dir / "1.png"
        if not base_file.exists():
            first_color = colors[0] if colors else ""
            first_prompt = f"{optimized_prompt}\nWhite background ecommerce product image, product centered.\nRequired product color/material element: {first_color}."
            generate_single_image(first_prompt, base_file, task.get("reference_image_path") or "")
        image_files.append(base_file)
        for index, color in enumerate(colors[1:], start=2):
            image_file = sku_dir / f"{index}.png"
            if not image_file.exists():
                if use_recolor:
                    recolor_image(base_file, image_file, color, task.get("sku") or "", colors[0] if colors else "红")
                else:
                    prompt = f"{optimized_prompt}\nWhite background ecommerce product image, product centered.\nRequired product color/material element: {color}."
                    generate_single_image(prompt, image_file, task.get("reference_image_path") or "")
            image_files.append(image_file)
        for index in range(len(image_files) + 1, 10):
            filler = sku_dir / f"{index}.png"
            if not filler.exists():
                shutil.copy2(base_file, filler)
            image_files.append(filler)
        save_task_fields(task, 执行状态="拼图中")
        compose_nine_grid(image_files[:9], main_file)
        final_file = main_file
        if str(task.get("add_logo") or "").strip().lower() in ("是", "yes", "y", "true", "1"):
            if not task.get("logo_path"):
                raise FileNotFoundError("Logo 文件未找到")
            add_logo_to_image(main_file, task.get("logo_path"), logo_file)
            final_file = logo_file
        ai_title = generate_ai_title(task)
        save_task_fields(
            task,
            执行状态="已完成",
            执行日志="",
            生成时间=now_iso(),
            输出文件地址=str(final_file),
            生成图片张数=len(image_files[:9]) + 1,
            图片颜色列表="/".join(colors),
            AI推荐标题=ai_title,
            优化后提示词=optimized_prompt,
        )
        return True
    except Exception as exc:
        message = openai_error_message(exc)
        status = "额度耗尽" if is_quota_error(message) else "失败"
        save_task_fields(task, 执行状态=status, 执行日志=message[:4000])
        return False


def stop_generation():
    queue_control["stop_requested"] = True
    queue_control["pause_requested"] = False
    queue_control["state"] = "stopped"
    for task in excel_load_tasks():
        if task["status"] == "generating":
            save_task_fields(task, 执行状态="待生成")
    return generate_status()


def generate_status():
    counts = {}
    for task in excel_load_tasks():
        counts[task["status"]] = counts.get(task["status"], 0) + 1
    pending = counts.get("pending", 0)
    generating = counts.get("generating", 0)
    total = sum(counts.values())
    completed = max(total - pending - generating, 0)
    state = queue_control["state"]
    if worker_running and not queue_control["pause_requested"] and not queue_control["stop_requested"]:
        state = "running"
    elif not worker_running and state == "running":
        state = "idle"
    return {"state": state, "running": worker_running, "paused": state == "paused", "pending": pending, "generating": generating, "completed": completed, "total": total}


def queue_all_paused():
    for task in excel_load_tasks():
        if task["status"] == "quota_exhausted":
            save_task_fields(task, 执行状态="待生成", 执行日志="")
    set_quota("unknown", "已加入待续跑队列，等待下一次 API 结果")
    start_worker()


def export_excel(table_mode="new", source_path=None, allow_source_fallback=True):
    path = active_excel_path()
    if path:
        return path
    EXPORT_DIR.mkdir(exist_ok=True)
    output_path = EXPORT_DIR / f"erp_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = IMAGE_GENERATION_SHEET
    ws.append(USER_EXCEL_COLUMNS + PROGRAM_EXCEL_COLUMNS)
    wb.save(output_path)
    return output_path


def app_state():
    counts = {}
    for task in excel_load_tasks():
        counts[task["status"]] = counts.get(task["status"], 0) + 1
=======
    if status in ("failed", "collaging_failed", "quality_failed", "rejected", "quota_exhausted"):
        return "失败"
    return "待生成"


def export_excel():
    EXPORT_DIR.mkdir(exist_ok=True)
    filename = f"erp_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = EXPORT_DIR / filename
    wb = Workbook()
    ws = wb.active
    ws.title = "ERP导出"
    columns = [
        "SPU",
        "图片生成提示词",
        "图片颜色",
        "图片命名",
        "是否添加品牌logo",
        "logo文件地址",
        "需校验品牌文字/元素",
        "优化后提示词",
        "提示词原文快照",
        "生成状态",
        "失败原因",
        "生成时间",
        "输出文件地址",
        "SKU",
    ]
    ws.append(columns)
    with connect() as conn:
        rows = conn.execute("SELECT * FROM tasks ORDER BY id ASC").fetchall()
    for row in rows:
        ws.append(
            [
                row["spu"] or "",
                row["prompt"] or "",
                row["color_name"] or "",
                row["image_name"] or "",
                row["add_logo"] or "",
                row["logo_path"] or "",
                row["brand_check"] or "",
                row["optimized_prompt"] or "",
                row["prompt_snapshot"] or "",
                export_generation_status(row["status"]),
                row["error_message"] or "",
                row["generated_at"] or "",
                row["image_path"] or "",
                row["sku"],
            ]
        )
    wb.save(path)
    return path


def app_state():
    with connect() as conn:
        counts = {
            row["status"]: row["count"]
            for row in conn.execute("SELECT status, COUNT(*) AS count FROM tasks GROUP BY status")
        }
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    return {
        "counts": counts,
        "quota": quota_state,
        "codex": api_config_state,
        "worker_running": worker_running,
        "generation": generate_status(),
<<<<<<< HEAD
        "detail_generation": detail_status(),
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        "statuses": STATUS_LABELS,
    }


HTTP_ROUTES = [
    ("GET", "/"),
    ("GET", "/docs"),
    ("GET", "/app.css"),
    ("GET", "/app.js"),
    ("GET", "/api/tasks"),
    ("GET", "/api/state"),
    ("GET", "/api/codex/status"),
    ("GET", "/api/codex/user-info"),
    ("GET", "/api/test-route"),
<<<<<<< HEAD
    ("GET", "/api/export/source-status"),
    ("GET", "/api/export"),
    ("POST", "/api/export/source"),
=======
    ("GET", "/api/export"),
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    ("GET", "/api/images/{sku}"),
    ("GET", "/output/{path}"),
    ("POST", "/api/download/batch"),
    ("GET", "/api/download/main/{sku}"),
    ("GET", "/api/download/all/{sku}"),
<<<<<<< HEAD
    ("GET", "/api/download/white/{sku}"),
    ("GET", "/api/download/detail/{sku}"),
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    ("GET", "/api/download/{sku}"),
    ("DELETE", "/api/tasks"),
    ("POST", "/api/manual"),
    ("POST", "/api/import"),
    ("POST", "/api/collage/test"),
    ("POST", "/api/collage/{sku:path}"),
    ("POST", "/api/generate/start"),
<<<<<<< HEAD
    ("POST", "/api/generate/detail"),
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    ("POST", "/api/generate/nine/{sku}"),
    ("POST", "/api/generate/pause"),
    ("POST", "/api/generate/stop"),
    ("GET", "/api/generate/status"),
    ("POST", "/api/quota/resume"),
    ("POST", "/api/codex/refresh"),
    ("POST", "/api/codex/set-config"),
    ("POST", "/api/codex/set-config/"),
    ("POST", "/api/codex/set-key"),
    ("POST", "/api/codex/set-key/"),
    ("POST", "/api/tasks/{task_id}/{action}"),
]


def docs_html():
    rows = "\n".join(
        f"<tr><td>{method}</td><td><code>{path}</code></td></tr>" for method, path in HTTP_ROUTES
    )
    return f"""
<!doctype html>
<html lang="zh-CN">
  <head>
    <meta charset="utf-8" />
    <title>SKU Workflow API Docs</title>
    <style>
      body {{ font-family: Segoe UI, Microsoft YaHei, sans-serif; margin: 32px; color: #1f2937; }}
      table {{ border-collapse: collapse; width: 100%; max-width: 920px; }}
      th, td {{ border-bottom: 1px solid #e5e7eb; padding: 10px 12px; text-align: left; }}
      th {{ background: #f8fafc; }}
      code {{ background: #f1f5f9; border-radius: 4px; padding: 2px 6px; }}
    </style>
  </head>
  <body>
    <h1>SKU Workflow API</h1>
    <p>当前环境未启用 FastAPI 自动文档，以下为本地服务已注册路由。</p>
    <table>
      <thead><tr><th>Method</th><th>Path</th></tr></thead>
      <tbody>{rows}</tbody>
    </table>
  </body>
</html>
""".strip()


def print_registered_routes():
    if FASTAPI_AVAILABLE and app is not None:
        for route in app.routes:
            methods = ", ".join(sorted(route.methods or []))
            print(f"Route: {route.path} [{methods}]", flush=True)
    else:
        for method, path in HTTP_ROUTES:
            print(f"Route: {path} [{method}]", flush=True)


def debug_print_registered_routes():
    print("已注册的路由:", flush=True)
    print_registered_routes()


app = FastAPI(title="SKU Workflow API", docs_url="/docs") if FASTAPI_AVAILABLE else None

if FASTAPI_AVAILABLE:
    @app.post("/api/collage/test")
    async def test_collage():
        return {"status": "ok"}

    @app.get("/api/test-route")
    async def test_route():
        return {"status": "ok"}

    @app.get("/")
    def fastapi_index():
        return FileResponse(STATIC_DIR / "index.html", media_type="text/html")

    @app.get("/app.css")
    def fastapi_css():
        return FileResponse(STATIC_DIR / "app.css", media_type="text/css")

    @app.get("/app.js")
    def fastapi_js():
        return FileResponse(STATIC_DIR / "app.js", media_type="application/javascript")

    @app.get("/api/tasks")
    def fastapi_tasks(status: str = "all", q: str = ""):
        return list_tasks(status, q)

    @app.get("/api/state")
    def fastapi_state():
        return app_state()

    @app.get("/api/codex/status")
    def fastapi_codex_status():
        load_api_base_url()
        return dict(api_config_state)

    @app.get("/api/codex/user-info")
    def fastapi_codex_user_info():
        return api_config_user_info()

    @app.post("/api/codex/set-key")
    @app.post("/api/codex/set-key/")
    @app.post("/api/codex/set-config")
    @app.post("/api/codex/set-config/")
    async def set_codex_config(request: FastAPIRequest):
        try:
            payload = await request.json()
            result = set_api_base_url(
                payload.get("api_base_url") or payload.get("api_key"),
                payload.get("access_token"),
            )
            return {"success": True, "message": "本地服务地址已保存", **result}
        except Exception as exc:
            return JSONResponse({"success": False, "error": str(exc)}, status_code=400)

    @app.post("/api/codex/refresh")
    def fastapi_codex_refresh():
        load_api_base_url()
        return dict(api_config_state)

    @app.get("/api/export")
<<<<<<< HEAD
    def fastapi_export(table: str = "new"):
        path = export_excel(table)
        return FileResponse(
            path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path.name,
        )

    @app.get("/api/export/source-status")
    def fastapi_export_source_status():
        return export_source_status()

    @app.post("/api/export/source")
    async def fastapi_export_with_source(request: FastAPIRequest, table: str = "new"):
        body = await request.body()
        if not body:
            return JSONResponse({"ok": False, "error": "请选择原始 Excel 文件"}, status_code=400)
        name = request.headers.get("X-Filename") or f"source_{uuid.uuid4().hex}.xlsx"
        safe_name = Path(name).name
        source_path = IMPORT_DIR / f"export_source_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_name}"
        source_path.write_bytes(body)
        set_config("last_import_path", str(source_path.resolve()))
        set_config("last_import_table_mode", table)
        path = export_excel(table, source_path=source_path, allow_source_fallback=False)
=======
    def fastapi_export():
        path = export_excel()
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        return FileResponse(
            path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path.name,
        )

    @app.get("/api/images/{sku:path}")
    def fastapi_images_sku(sku: str):
        task = get_task_by_sku(sku)
        if not task:
            return JSONResponse({"ok": False, "error": "Task not found"}, status_code=404)
        images = image_list_for_sku(sku)
        print(f"DEBUG images route hit: sku={sku}, count={len(images)}", flush=True)
        return {"ok": True, "sku": sku, "images": images}

    @app.post("/api/download/batch")
    async def fastapi_download_batch(request: FastAPIRequest):
        try:
            payload = await request.json()
            path = zip_images_for_skus(payload.get("skus") or [])
            return FileResponse(path, media_type="application/zip", filename=path.name)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.get("/api/download/main/{sku:path}")
    def fastapi_download_main_sku(sku: str):
        task = get_task_by_sku(sku)
        image_path = download_image_path_for_sku(sku)
        print(f"DEBUG download main route hit: sku={sku}, image_path={image_path}", flush=True)
        if not image_path:
            return JSONResponse({"ok": False, "error": "图片未生成"}, status_code=404)
        return FileResponse(image_path, media_type="image/png", filename=f"{download_base_name_for_task(task, sku)}_主图.png")

    @app.get("/api/download/all/{sku:path}")
    def fastapi_download_all_sku(sku: str):
        try:
            zip_path = zip_all_images_for_sku(sku)
            print(f"DEBUG download all route hit: sku={sku}, zip_path={zip_path}", flush=True)
            return FileResponse(zip_path, media_type="application/zip", filename=zip_path.name)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)

<<<<<<< HEAD
    @app.get("/api/download/white/{sku:path}")
    def fastapi_download_white_sku(sku: str):
        try:
            zip_path = zip_all_images_for_sku(sku, "white")
            return FileResponse(zip_path, media_type="application/zip", filename=zip_path.name)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)

    @app.get("/api/download/detail/{sku:path}")
    def fastapi_download_detail_sku(sku: str):
        try:
            zip_path = zip_all_images_for_sku(sku, "detail")
            return FileResponse(zip_path, media_type="application/zip", filename=zip_path.name)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)

=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    @app.get("/api/download/{sku:path}")
    def fastapi_download_sku(sku: str):
        try:
            zip_path = zip_all_images_for_sku(sku)
            print(f"DEBUG download route hit: sku={sku}, zip_path={zip_path}", flush=True)
            return FileResponse(zip_path, media_type="application/zip", filename=zip_path.name)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=404)

    @app.delete("/api/tasks")
    async def fastapi_delete_tasks(request: FastAPIRequest):
        try:
            payload = await request.json()
            return delete_tasks_by_skus(payload.get("skus") or [])
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.get("/output/{file_path:path}")
    def fastapi_output(file_path: str):
        requested = (OUTPUT_DIR / file_path).resolve()
        if OUTPUT_DIR.resolve() not in requested.parents and requested != OUTPUT_DIR.resolve():
            return JSONResponse({"error": "Forbidden"}, status_code=403)
        if not requested.exists():
            return JSONResponse({"error": "Not found"}, status_code=404)
        return FileResponse(requested, media_type="image/png")

    @app.post("/api/manual")
    async def fastapi_manual(request: FastAPIRequest):
        try:
            add_manual_task(await request.json())
            if load_api_base_url():
                start_worker()
            return {"ok": True}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.post("/api/import")
<<<<<<< HEAD
    async def fastapi_import(request: FastAPIRequest, table: str = "auto"):
=======
    async def fastapi_import(request: FastAPIRequest):
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        try:
            body = await request.body()
            name = request.headers.get("X-Filename") or f"upload_{uuid.uuid4().hex}.xlsx"
            safe_name = Path(name).name
            path = IMPORT_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_name}"
            path.write_bytes(body)
<<<<<<< HEAD
            stats = import_excel(path, table)
            if (stats.get("imported") or stats.get("updated")) and load_api_base_url():
=======
            stats = import_excel(path)
            if stats.get("imported") and load_api_base_url():
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
                start_worker()
            return {"ok": True, "stats": stats}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.post("/api/generate/start")
    def fastapi_generate_start():
        try:
            require_api_base_url()
            start_worker()
            return {"ok": True}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

<<<<<<< HEAD
    @app.post("/api/generate/detail")
    async def fastapi_generate_detail(request: FastAPIRequest):
        try:
            payload = await request.json()
            require_api_base_url()
            return start_detail_generation(payload.get("mode", "all"), payload.get("skus") or [])
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    @app.post("/api/collage/{sku:path}")
    def fastapi_collage_sku(sku: str):
        try:
            return rerun_collage_for_sku(sku)
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.post("/api/generate/nine/{sku:path}")
    def fastapi_generate_nine(sku: str):
        try:
            task = get_task_by_sku(sku)
            if not task:
                return JSONResponse({"ok": False, "error": "Task not found"}, status_code=404)
            reset_for_generation(task["id"])
            require_api_base_url()
            start_worker()
            return {"ok": True}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.post("/api/generate/pause")
    def fastapi_generate_pause():
        return pause_generation()

    @app.post("/api/generate/stop")
    def fastapi_generate_stop():
        return stop_generation()

    @app.get("/api/generate/status")
    def fastapi_generate_status():
        return generate_status()

    @app.post("/api/quota/resume")
    def fastapi_quota_resume():
        try:
            require_api_base_url()
            queue_all_paused()
            return {"ok": True}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)

    @app.post("/api/tasks/{task_id}/{action}")
    async def fastapi_task_action(task_id: int, action: str, request: FastAPIRequest):
        try:
            if action == "approve":
                update_status(task_id, "approved")
            elif action == "reject":
                update_status(task_id, "rejected")
            elif action == "quality-fail":
                payload = await request.json()
                update_status(task_id, "quality_failed", payload.get("message", "质检失败"))
            elif action == "retry":
                require_api_base_url()
                reset_for_generation(task_id)
                start_worker()
            elif action == "queue":
                reset_for_generation(task_id)
            else:
                return JSONResponse({"ok": False, "error": "Not found"}, status_code=404)
            return {"ok": True}
        except Exception as exc:
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=400)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        LOG_DIR.mkdir(exist_ok=True)
        with (LOG_DIR / "server.log").open("a", encoding="utf-8") as fh:
            fh.write(f"{now_iso()} {self.address_string()} {fmt % args}\n")

<<<<<<< HEAD
    def write_body(self, body):
        try:
            self.wfile.write(body)
        except (ConnectionAbortedError, BrokenPipeError):
            pass

=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    def send_json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
<<<<<<< HEAD
        self.write_body(body)
=======
        self.wfile.write(body)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e

    def send_file(self, path, content_type="application/octet-stream", download_name=None):
        if not path.exists():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        body = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        if download_name:
            ascii_name = "".join(ch if ord(ch) < 128 else "_" for ch in download_name)
            encoded_name = quote(download_name)
            self.send_header(
                "Content-Disposition",
                f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded_name}",
            )
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
<<<<<<< HEAD
        self.write_body(body)
=======
        self.wfile.write(body)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0"))
        return json.loads(self.rfile.read(length).decode("utf-8") or "{}")

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.send_file(STATIC_DIR / "index.html", "text/html; charset=utf-8")
        elif parsed.path == "/docs":
            body = docs_html().encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
<<<<<<< HEAD
            self.write_body(body)
=======
            self.wfile.write(body)
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        elif parsed.path == "/app.css":
            self.send_file(STATIC_DIR / "app.css", "text/css; charset=utf-8")
        elif parsed.path == "/app.js":
            self.send_file(STATIC_DIR / "app.js", "application/javascript; charset=utf-8")
        elif parsed.path == "/api/tasks":
            qs = parse_qs(parsed.query)
            self.send_json(list_tasks(qs.get("status", ["all"])[0], qs.get("q", [""])[0]))
        elif parsed.path == "/api/state":
            self.send_json(app_state())
        elif parsed.path == "/api/generate/status":
            self.send_json(generate_status())
        elif parsed.path == "/api/codex/status":
            load_api_base_url()
            self.send_json(dict(api_config_state))
        elif parsed.path == "/api/codex/user-info":
            self.send_json(api_config_user_info())
        elif parsed.path == "/api/test-route":
            self.send_json({"status": "ok"})
<<<<<<< HEAD
        elif parsed.path == "/api/export/source-status":
            self.send_json(export_source_status())
        elif parsed.path == "/api/export":
            qs = parse_qs(parsed.query)
            path = export_excel(qs.get("table", ["new"])[0])
=======
        elif parsed.path == "/api/export":
            path = export_excel()
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            self.send_file(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", path.name)
        elif parsed.path.startswith("/api/images/"):
            sku = unquote(parsed.path.split("/api/images/", 1)[1])
            task = get_task_by_sku(sku)
            if not task:
                self.send_json({"ok": False, "error": "Task not found"}, 404)
                return
            images = image_list_for_sku(sku)
            print(f"DEBUG images route hit: sku={sku}, count={len(images)}", flush=True)
            self.send_json({"ok": True, "sku": sku, "images": images})
        elif parsed.path.startswith("/api/download/main/"):
            sku = unquote(parsed.path.split("/api/download/main/", 1)[1])
            task = get_task_by_sku(sku)
            image_path = download_image_path_for_sku(sku)
            print(f"DEBUG download main route hit: sku={sku}, image_path={image_path}", flush=True)
            if not image_path:
                self.send_json({"ok": False, "error": "图片未生成"}, 404)
                return
            self.send_file(image_path, "image/png", f"{download_base_name_for_task(task, sku)}_主图.png")
        elif parsed.path.startswith("/api/download/all/"):
            sku = unquote(parsed.path.split("/api/download/all/", 1)[1])
            zip_path = zip_all_images_for_sku(sku)
            print(f"DEBUG download all route hit: sku={sku}, zip_path={zip_path}", flush=True)
            self.send_file(zip_path, "application/zip", zip_path.name)
<<<<<<< HEAD
        elif parsed.path.startswith("/api/download/white/"):
            sku = unquote(parsed.path.split("/api/download/white/", 1)[1])
            zip_path = zip_all_images_for_sku(sku, "white")
            self.send_file(zip_path, "application/zip", zip_path.name)
        elif parsed.path.startswith("/api/download/detail/"):
            sku = unquote(parsed.path.split("/api/download/detail/", 1)[1])
            zip_path = zip_all_images_for_sku(sku, "detail")
            self.send_file(zip_path, "application/zip", zip_path.name)
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        elif parsed.path.startswith("/api/download/"):
            sku = unquote(parsed.path.split("/api/download/", 1)[1])
            zip_path = zip_all_images_for_sku(sku)
            print(f"DEBUG download route hit: sku={sku}, zip_path={zip_path}", flush=True)
            self.send_file(zip_path, "application/zip", zip_path.name)
        elif parsed.path.startswith("/output/"):
            relative_output_path = unquote(parsed.path.split("/output/", 1)[1])
            requested = (OUTPUT_DIR / relative_output_path).resolve()
            if OUTPUT_DIR.resolve() not in requested.parents and requested != OUTPUT_DIR.resolve():
                self.send_error(HTTPStatus.FORBIDDEN)
                return
            self.send_file(requested, "image/png")
        else:
            self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self):
        try:
            parsed = urlparse(self.path)
            normalized_path = parsed.path.rstrip("/") or "/"
            if normalized_path in ("/api/codex/set-config", "/api/codex/set-key"):
                print(f"DEBUG: set local API config called: {parsed.path}", flush=True)
                payload = self.read_json()
                self.send_json(
                    set_api_base_url(
                        payload.get("api_base_url") or payload.get("api_key"),
                        payload.get("access_token"),
                    )
                )
            elif parsed.path == "/api/manual":
                add_manual_task(self.read_json())
                if load_api_base_url():
                    start_worker()
                self.send_json({"ok": True})
            elif parsed.path == "/api/import":
                self.handle_import()
<<<<<<< HEAD
            elif parsed.path == "/api/export/source":
                self.handle_export_source()
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            elif parsed.path == "/api/generate/start":
                require_api_base_url()
                start_worker()
                self.send_json({"ok": True})
<<<<<<< HEAD
            elif parsed.path == "/api/generate/detail":
                payload = self.read_json()
                require_api_base_url()
                self.send_json(start_detail_generation(payload.get("mode", "all"), payload.get("skus") or []))
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
            elif parsed.path == "/api/collage/test":
                self.send_json({"status": "ok"})
            elif parsed.path.startswith("/api/collage/"):
                sku = unquote(parsed.path.split("/api/collage/", 1)[1])
                self.send_json(rerun_collage_for_sku(sku))
            elif parsed.path.startswith("/api/generate/nine/"):
                sku = unquote(parsed.path.split("/api/generate/nine/", 1)[1])
                task = get_task_by_sku(sku)
                if not task:
                    self.send_json({"ok": False, "error": "Task not found"}, 404)
                    return
                reset_for_generation(task["id"])
                require_api_base_url()
                start_worker()
                self.send_json({"ok": True})
            elif parsed.path == "/api/generate/pause":
                self.send_json(pause_generation())
            elif parsed.path == "/api/generate/stop":
                self.send_json(stop_generation())
            elif parsed.path == "/api/quota/resume":
                require_api_base_url()
                queue_all_paused()
                self.send_json({"ok": True})
            elif parsed.path == "/api/codex/refresh":
                load_api_base_url()
                self.send_json(dict(api_config_state))
            elif parsed.path == "/api/download/batch":
                payload = self.read_json()
                path = zip_images_for_skus(payload.get("skus") or [])
                self.send_file(path, "application/zip", path.name)
            elif parsed.path.startswith("/api/tasks/"):
                self.handle_task_action(parsed.path)
            else:
                self.send_error(HTTPStatus.NOT_FOUND)
        except Exception as exc:
            traceback.print_exc()
            self.send_json({"ok": False, "error": str(exc)}, 400)

    def do_DELETE(self):
        try:
            parsed = urlparse(self.path)
            if parsed.path == "/api/tasks":
                payload = self.read_json()
                self.send_json(delete_tasks_by_skus(payload.get("skus") or []))
            else:
                self.send_error(HTTPStatus.NOT_FOUND)
        except Exception as exc:
            traceback.print_exc()
            self.send_json({"ok": False, "error": str(exc)}, 400)

    def handle_import(self):
<<<<<<< HEAD
        parsed = urlparse(self.path)
        qs = parse_qs(parsed.query)
        table_mode = qs.get("table", ["auto"])[0]
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
        length = int(self.headers.get("Content-Length", "0"))
        name = self.headers.get("X-Filename") or f"upload_{uuid.uuid4().hex}.xlsx"
        safe_name = Path(name).name
        path = IMPORT_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_name}"
        path.write_bytes(self.rfile.read(length))
<<<<<<< HEAD
        stats = import_excel(path, table_mode)
        if (stats.get("imported") or stats.get("updated")) and load_api_base_url():
            start_worker()
        self.send_json({"ok": True, "stats": stats})

    def handle_export_source(self):
        parsed = urlparse(self.path)
        qs = parse_qs(parsed.query)
        table_mode = qs.get("table", ["new"])[0]
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0:
            self.send_json({"ok": False, "error": "请选择原始 Excel 文件"}, 400)
            return
        name = self.headers.get("X-Filename") or f"source_{uuid.uuid4().hex}.xlsx"
        safe_name = Path(name).name
        source_path = IMPORT_DIR / f"export_source_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_name}"
        source_path.write_bytes(self.rfile.read(length))
        set_config("last_import_path", str(source_path.resolve()))
        set_config("last_import_table_mode", table_mode)
        path = export_excel(table_mode, source_path=source_path, allow_source_fallback=False)
        self.send_file(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", path.name)

=======
        stats = import_excel(path)
        if stats.get("imported") and load_api_base_url():
            start_worker()
        self.send_json({"ok": True, "stats": stats})

>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    def handle_task_action(self, path):
        parts = path.strip("/").split("/")
        if len(parts) != 4:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        task_id = int(parts[2])
        action = parts[3]
        if action == "approve":
            update_status(task_id, "approved")
        elif action == "reject":
            update_status(task_id, "rejected")
        elif action == "quality-fail":
            payload = self.read_json()
            update_status(task_id, "quality_failed", payload.get("message", "质检失败"))
        elif action == "retry":
            require_api_base_url()
            reset_for_generation(task_id)
            start_worker()
        elif action == "queue":
            reset_for_generation(task_id)
        else:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        self.send_json({"ok": True})


def main():
<<<<<<< HEAD
    parser = argparse.ArgumentParser(add_help=True)
    parser.add_argument("--table", choices=["auto", "old", "new"], default="auto", help="Excel 表结构：old=图片生成模板，new=跨境电商统一运营模板")
    parser.add_argument("--import-excel", help="导入指定 Excel 后退出，不启动网页服务")
    args, _ = parser.parse_known_args()
    if args.import_excel:
        init_db()
        set_config("table_mode", args.table)
        stats = import_excel(Path(args.import_excel), args.table)
        print(json.dumps(stats, ensure_ascii=False), flush=True)
        return
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    print("Server started, waiting for requests...", flush=True)
    print(f"DEBUG app.py path: {Path(__file__).resolve()}", flush=True)
    print(f"DEBUG cwd: {Path.cwd()}", flush=True)
    print(f"DEBUG server mode: {'FastAPI' if FASTAPI_AVAILABLE else 'fallback-http'}", flush=True)
    ensure_admin()
    init_db()
<<<<<<< HEAD
    set_config("table_mode", args.table)
=======
>>>>>>> 0afc0ab16f9d6e1a79023e8edd542d3d99d93f8e
    api_base_url = load_api_base_url()
    print_local_api_config()
    if not api_base_url:
        print("WARNING: 未配置 Codex 本地服务地址：请先在网页右上角配置本地服务", flush=True)
    host = "127.0.0.1"
    port = int(os.environ.get("SKU_APP_PORT", "8765"))
    print_registered_routes()
    if FASTAPI_AVAILABLE:
        print(f"SKU workflow dashboard running at http://{host}:{port}", flush=True)
        uvicorn.run(app, host=host, port=port, log_level="info")
        return
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"SKU workflow dashboard running at http://{host}:{port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
